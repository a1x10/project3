"""Что Стелла делает сама: стикеры и картинки (рисует в SVG), мемы, документы, презентации, таблицы, видео.

Всё рисуется и собирается здесь, на сервере, из того, что написала модель: никакой чужой код не выполняется.
SVG чистится (без скриптов, внешних ссылок и вставленных файлов), видео собирает ffmpeg с фиксированными
параметрами, в таблицах формулы превращаются в обычный текст.
"""

from __future__ import annotations

import asyncio
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from typing import Any

from PIL import Image, ImageDraw, ImageFont

MAX_SVG_CHARS = 200_000
MAX_SVG_TAGS = 4000
STICKER_SIZE = 512
STICKER_MAX_BYTES = 100_000  # WhatsApp не принимает статичный стикер больше 100 КБ
VIDEO_FPS = 24
VIDEO_MAX_SECONDS = 40
VIDEO_MAX_SCENES = 12
FONT_DIRS = ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "/usr/share/fonts/TTF")

SIZES = {"square": (1080, 1080), "wide": (1280, 720), "tall": (720, 1280)}
VIDEO_SIZES = {"square": (720, 720), "wide": (1280, 720), "tall": (720, 1280)}

MIME = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "webp": "image/webp",
    "mp4": "video/mp4",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class MediaError(Exception):
    """Понятная модели ошибка: что не так с тем, что она прислала."""


@dataclass
class MediaFile:
    kind: str  # image | video | document | sticker
    data: bytes
    mimetype: str
    filename: str


# ---------- шрифты ----------
def font_path(bold: bool = False, mono: bool = False) -> str | None:
    name = "DejaVuSansMono" if mono else "DejaVuSans"
    name += "-Bold.ttf" if bold else ".ttf"
    for d in FONT_DIRS:
        p = os.path.join(d, name)
        if os.path.exists(p):
            return p
    return None


def _font(size: int, bold: bool = True) -> ImageFont.ImageFont:
    path = font_path(bold)
    return ImageFont.truetype(path, size) if path else ImageFont.load_default()


# ---------- SVG ----------
_DROP_BLOCKS = re.compile(r"<(script|foreignObject|iframe|object|embed|audio|video|image)\b.*?(</\1\s*>|/>)",
                          re.IGNORECASE | re.DOTALL)
_DROP_TAGS = re.compile(r"<(script|foreignObject|iframe|object|embed|audio|video|image)\b[^>]*>", re.IGNORECASE)
_EVENT_ATTR = re.compile(r"\son[a-z]+\s*=\s*(\"[^\"]*\"|'[^']*')", re.IGNORECASE)
_HREF = re.compile(r"\s(xlink:)?href\s*=\s*(\"([^\"]*)\"|'([^']*)')", re.IGNORECASE)
_CSS_URL = re.compile(r"url\(\s*['\"]?(?!#)[^)]*\)", re.IGNORECASE)
_CSS_IMPORT = re.compile(r"@import[^;]*;?", re.IGNORECASE)


def sanitize_svg(svg: str) -> str:
    """Оставляем только рисунок: без скриптов, внешних ссылок, вставленных картинок и сущностей XML."""
    if not isinstance(svg, str) or "<svg" not in svg.lower():
        raise MediaError("нужен SVG-код: <svg ...>...</svg>")
    if len(svg) > MAX_SVG_CHARS:
        raise MediaError("SVG слишком большой, нарисуй проще")
    if re.search(r"<!(DOCTYPE|ENTITY)", svg, re.IGNORECASE):
        raise MediaError("без DOCTYPE и ENTITY")
    start = svg.lower().index("<svg")
    end = svg.lower().rfind("</svg>")
    svg = svg[start:end + 6] if end > start else svg[start:]
    svg = _DROP_BLOCKS.sub("", svg)
    svg = _DROP_TAGS.sub("", svg)
    svg = _EVENT_ATTR.sub("", svg)
    svg = _HREF.sub(lambda m: m.group(0) if (m.group(3) or m.group(4) or "").startswith("#") else "", svg)
    svg = _CSS_IMPORT.sub("", svg)
    svg = _CSS_URL.sub("none", svg)
    if svg.count("<") > MAX_SVG_TAGS:
        raise MediaError("слишком много элементов, нарисуй проще")
    if not re.match(r"<svg\b[^>]*\sxmlns\s*=", svg, re.IGNORECASE):
        svg = svg.replace("<svg", '<svg xmlns="http://www.w3.org/2000/svg"', 1)
    return svg


# Рисуем в отдельном процессе с лимитом времени и памяти: хитрый SVG (вложенные <use>, узоры, маски)
# может раздуться в миллионы элементов — тогда процесс просто убивается, а бот работает дальше.
_SVG_WORKER = r"""
import json, resource, sys
resource.setrlimit(resource.RLIMIT_AS, (1536 * 1024 * 1024, 1536 * 1024 * 1024))
resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
import cairosvg
req = json.loads(sys.stdin.buffer.read())
# unsafe=False: без внешних файлов и XML-сущностей (cairosvg сам берёт безопасный загрузчик)
png = cairosvg.svg2png(bytestring=req["svg"].encode("utf-8"), output_width=req["w"], output_height=req["h"],
                       unsafe=False)
sys.stdout.buffer.write(png)
"""
SVG_RENDER_TIMEOUT = 20
MAX_RENDER_PIXELS = 2200 * 2200


def render_svg(svg: str, width: int, height: int) -> Image.Image:
    clean = sanitize_svg(svg)
    if width * height > MAX_RENDER_PIXELS:
        raise MediaError("слишком большой размер картинки")
    try:
        proc = subprocess.run([sys.executable, "-I", "-c", _SVG_WORKER],
                              input=json.dumps({"svg": clean, "w": width, "h": height}).encode(),
                              capture_output=True, timeout=SVG_RENDER_TIMEOUT)
    except subprocess.TimeoutExpired as exc:
        raise MediaError("рисунок слишком сложный, нарисуй проще") from exc
    if proc.returncode != 0 or not proc.stdout:
        err = (proc.stderr or b"").decode(errors="replace").strip().splitlines()
        reason = err[-1][:200] if err else f"код {proc.returncode}"
        if proc.returncode < 0 or "MemoryError" in reason:
            reason = "слишком сложный рисунок"
        raise MediaError(f"не получилось нарисовать SVG: {reason}")
    try:
        return Image.open(io.BytesIO(proc.stdout)).convert("RGBA")
    except Exception as exc:
        raise MediaError("не получилось нарисовать SVG") from exc


# ---------- картинка для модели (она «видит» фото и стикеры) ----------
MODEL_IMAGE_SIDE = 1568  # больше модели не нужно — только дороже


def image_for_model(data: bytes) -> tuple[str, bytes] | None:
    """Фото/стикер из WhatsApp -> (media_type, байты) для модели: первый кадр, не больше 1568 px.
    None — если это не картинка."""
    try:
        img = Image.open(io.BytesIO(data))
        img.seek(0)
        img.load()
    except Exception:
        return None
    if img.width * img.height > 50_000_000:
        return None
    has_alpha = img.mode in ("RGBA", "LA", "P") and "transparency" in img.info or img.mode in ("RGBA", "LA")
    img = img.convert("RGBA" if has_alpha else "RGB")
    img.thumbnail((MODEL_IMAGE_SIDE, MODEL_IMAGE_SIDE), Image.LANCZOS)
    buf = io.BytesIO()
    if has_alpha:  # стикеры с прозрачностью — на белом фоне, чтобы модель видела контуры
        bg = Image.new("RGB", img.size, (255, 255, 255))
        bg.paste(img, mask=img.split()[-1])
        img = bg
    img.save(buf, "JPEG", quality=85)
    return "image/jpeg", buf.getvalue()


# ---------- стикеры и картинки ----------
def to_sticker(img: Image.Image) -> bytes:
    """Квадрат 512×512 с прозрачным фоном, WebP до 100 КБ."""
    img = img.convert("RGBA")
    img.thumbnail((STICKER_SIZE, STICKER_SIZE), Image.LANCZOS)
    canvas = Image.new("RGBA", (STICKER_SIZE, STICKER_SIZE), (0, 0, 0, 0))
    canvas.paste(img, ((STICKER_SIZE - img.width) // 2, (STICKER_SIZE - img.height) // 2), img)
    for quality in (90, 75, 60, 45, 30):
        buf = io.BytesIO()
        canvas.save(buf, "WEBP", quality=quality, method=6)
        if buf.tell() <= STICKER_MAX_BYTES:
            return buf.getvalue()
    raise MediaError("стикер получился слишком тяжёлым, нарисуй проще")


def sticker_from_svg(svg: str) -> MediaFile:
    return MediaFile("sticker", to_sticker(render_svg(svg, STICKER_SIZE, STICKER_SIZE)), MIME["webp"], "sticker.webp")


def picture_from_svg(svg: str, shape: str = "square") -> MediaFile:
    w, h = SIZES.get(shape, SIZES["square"])
    img = render_svg(svg, w, h)
    bg = Image.new("RGB", img.size, (255, 255, 255))
    bg.paste(img, mask=img.split()[-1])
    buf = io.BytesIO()
    bg.save(buf, "PNG", optimize=True)
    return MediaFile("image", buf.getvalue(), MIME["png"], "picture.png")


def _wrap(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_w: int) -> list[str]:
    lines: list[str] = []
    for para in text.split("\n"):
        words, line = para.split(), ""
        for w in words:
            probe = f"{line} {w}".strip()
            if draw.textlength(probe, font=font) <= max_w or not line:
                line = probe
            else:
                lines.append(line)
                line = w
        lines.append(line)
    return [ln for ln in lines if ln] or [""]


def _caption(img: Image.Image, text: str, where: str) -> None:
    """Белый текст с чёрной обводкой — как в мемах."""
    if not text.strip():
        return
    draw = ImageDraw.Draw(img)
    size = max(18, img.width // 9)
    while True:
        font = _font(size)
        lines = _wrap(draw, text.upper(), font, int(img.width * 0.92))
        line_h = int(size * 1.15)
        if (len(lines) * line_h < img.height * 0.4 and max(draw.textlength(ln, font=font) for ln in lines)
                <= img.width * 0.95) or size <= 14:
            break
        size -= 2
    total = len(lines) * line_h
    y = int(img.height * 0.03) if where == "top" else img.height - total - int(img.height * 0.04)
    for ln in lines:
        x = (img.width - draw.textlength(ln, font=font)) / 2
        draw.text((x, y), ln, font=font, fill="white", stroke_width=max(2, size // 12), stroke_fill="black")
        y += line_h


def meme(image: bytes, top: str = "", bottom: str = "", as_sticker: bool = False) -> MediaFile:
    try:
        img = Image.open(io.BytesIO(image))
        img.seek(0)
        img = img.convert("RGBA")
    except Exception as exc:
        raise MediaError("не получилось открыть картинку") from exc
    img.thumbnail((1024, 1024), Image.LANCZOS)
    _caption(img, top[:120], "top")
    _caption(img, bottom[:120], "bottom")
    if as_sticker:
        return MediaFile("sticker", to_sticker(img), MIME["webp"], "meme.webp")
    bg = Image.new("RGB", img.size, (255, 255, 255))
    bg.paste(img, mask=img.split()[-1])
    buf = io.BytesIO()
    bg.save(buf, "JPEG", quality=88)
    return MediaFile("image", buf.getvalue(), MIME["jpg"], "meme.jpg")


# ---------- документы ----------
def _safe_name(title: str, ext: str) -> str:
    base = re.sub(r"[^\w\- ]+", "", title or "", flags=re.UNICODE).strip()[:60] or "file"
    return f"{base}.{ext}"


def _blocks(blocks: Any) -> list[dict[str, Any]]:
    if not isinstance(blocks, list) or not blocks:
        raise MediaError("blocks: список блоков (heading, paragraph, bullets, numbered, table, quote)")
    out = []
    for b in blocks[:400]:
        if isinstance(b, str):
            b = {"type": "paragraph", "text": b}
        if isinstance(b, dict):
            out.append(b)
    return out


def _cell(v: Any) -> str:
    return "" if v is None else str(v)[:500]


def document_docx(title: str, blocks: Any) -> MediaFile:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    if title:
        doc.add_heading(title[:200], level=0)
    for b in _blocks(blocks):
        kind = b.get("type", "paragraph")
        text = str(b.get("text") or "")[:5000]
        if kind == "heading":
            doc.add_heading(text, level=min(max(int(b.get("level") or 1), 1), 3))
        elif kind in ("bullets", "numbered"):
            style = "List Bullet" if kind == "bullets" else "List Number"
            for item in (b.get("items") or [])[:100]:
                doc.add_paragraph(str(item)[:1000], style=style)
        elif kind == "table":
            rows = [r for r in (b.get("rows") or []) if isinstance(r, list)][:200]
            if rows:
                cols = min(max(len(r) for r in rows), 12)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for i, r in enumerate(rows):
                    for j in range(cols):
                        table.cell(i, j).text = _cell(r[j] if j < len(r) else "")
                        if i == 0:
                            for run in table.cell(i, j).paragraphs[0].runs:
                                run.bold = True
        elif kind == "quote":
            doc.add_paragraph(text, style="Intense Quote")
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return MediaFile("document", buf.getvalue(), MIME["docx"], _safe_name(title, "docx"))


def document_pdf(title: str, blocks: Any) -> MediaFile:
    from fpdf import FPDF

    regular, bold = font_path(False), font_path(True)
    if not regular or not bold:
        raise MediaError("на сервере нет шрифта для PDF")
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("DejaVu", "", regular)
    pdf.add_font("DejaVu", "B", bold)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def para(text: str, size: int = 12, style: str = "", gap: float = 2) -> None:
        pdf.set_font("DejaVu", style, size)
        pdf.multi_cell(width, size * 0.55, text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    if title:
        para(title[:200], 20, "B", 4)
    for b in _blocks(blocks):
        kind = b.get("type", "paragraph")
        text = str(b.get("text") or "")[:5000]
        if kind == "heading":
            para(text, 16 if int(b.get("level") or 1) <= 1 else 14, "B", 2)
        elif kind in ("bullets", "numbered"):
            for i, item in enumerate((b.get("items") or [])[:100], 1):
                para(f"{'•' if kind == 'bullets' else f'{i}.'} {str(item)[:1000]}", 12, "", 1)
            pdf.ln(2)
        elif kind == "table":
            rows = [[_cell(c) for c in r][:8] for r in (b.get("rows") or []) if isinstance(r, list)][:200]
            if rows:
                pdf.set_font("DejaVu", "", 10)
                with pdf.table(first_row_as_headings=True) as table:
                    cols = max(len(r) for r in rows)
                    for r in rows:
                        row = table.row()
                        for j in range(cols):
                            row.cell(r[j] if j < len(r) else "")
                pdf.ln(3)
        elif kind == "quote":
            para(f"«{text}»", 12, "", 2)
        else:
            para(text)
    return MediaFile("document", bytes(pdf.output()), MIME["pdf"], _safe_name(title, "pdf"))


def presentation(title: str, subtitle: str, slides: Any) -> MediaFile:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    if not isinstance(slides, list) or not slides:
        raise MediaError("slides: список слайдов {title, bullets[], notes?, svg?}")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    dark, accent, light = RGBColor(0x1E, 0x1B, 0x2E), RGBColor(0xFF, 0x7A, 0x59), RGBColor(0xF5, 0xF3, 0xFF)

    def paint(slide) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = dark

    def text_box(slide, x, y, w, h, text, size, color, bold=False):
        box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = text
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
        return tf

    blank = prs.slide_layouts[6]
    first = prs.slides.add_slide(blank)
    paint(first)
    text_box(first, 700000, 2200000, 10800000, 1500000, (title or "Презентация")[:150], 44, light, True)
    if subtitle:
        text_box(first, 700000, 3700000, 10800000, 900000, subtitle[:200], 22, accent)
    for s in slides[:30]:
        if not isinstance(s, dict):
            continue
        slide = prs.slides.add_slide(blank)
        paint(slide)
        text_box(slide, 600000, 400000, 11000000, 1000000, str(s.get("title") or "")[:150], 32, accent, True)
        bullets = [str(b)[:300] for b in (s.get("bullets") or []) if str(b).strip()][:10]
        picture = None
        if s.get("svg"):
            img = render_svg(str(s["svg"]), 900, 900)
            buf = io.BytesIO()
            img.save(buf, "PNG")
            buf.seek(0)
            picture = buf
        text_w = 6200000 if picture else 11000000
        if bullets:
            tf = text_box(slide, 600000, 1600000, text_w, 4800000, "• " + bullets[0], 22, light)
            for b in bullets[1:]:
                p = tf.add_paragraph()
                p.text = "• " + b
                p.space_before = Pt(10)
                for r in p.runs:
                    r.font.size, r.font.color.rgb = Pt(22), light
        if picture:
            slide.shapes.add_picture(picture, Emu(7200000), Emu(1500000), height=Emu(4500000))
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])[:3000]
    buf = io.BytesIO()
    prs.save(buf)
    return MediaFile("document", buf.getvalue(), MIME["pptx"], _safe_name(title, "pptx"))


def spreadsheet(title: str, sheets: Any) -> MediaFile:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    if not isinstance(sheets, list) or not sheets:
        raise MediaError("sheets: список листов {name, rows: [[...], ...]}")
    wb = Workbook()
    wb.remove(wb.active)
    for i, sh in enumerate(sheets[:10]):
        if not isinstance(sh, dict):
            continue
        name = re.sub(r"[\[\]:*?/\\]", "", str(sh.get("name") or f"Лист{i + 1}"))[:31] or f"Лист{i + 1}"
        ws = wb.create_sheet(name)
        widths: dict[int, int] = {}
        for r, row in enumerate((sh.get("rows") or [])[:2000], 1):
            if not isinstance(row, list):
                continue
            for c, value in enumerate(row[:50], 1):
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell = ws.cell(r, c, value)
                else:
                    text = _cell(value)
                    cell = ws.cell(r, c, text)
                    cell.data_type = "s"  # «=…» — просто текст, а не формула
                if r == 1:
                    cell.font = Font(bold=True)
                widths[c] = max(widths.get(c, 0), len(str(value or "")))
        for c, w in widths.items():
            ws.column_dimensions[ws.cell(1, c).column_letter].width = min(max(w + 2, 8), 60)
    if not wb.sheetnames:
        raise MediaError("пустая таблица")
    buf = io.BytesIO()
    wb.save(buf)
    return MediaFile("document", buf.getvalue(), MIME["xlsx"], _safe_name(title, "xlsx"))


# ---------- видео ----------
def _scene_card(scene: dict[str, Any], size: tuple[int, int]) -> Image.Image:
    w, h = size
    if scene.get("svg"):
        img = render_svg(str(scene["svg"]), w, h)
        bg = Image.new("RGBA", (w, h), _color(scene.get("bg"), (20, 18, 32)) + (255,))
        bg.alpha_composite(img)
        img = bg
    else:
        img = Image.new("RGBA", (w, h), _color(scene.get("bg"), (30, 27, 46)) + (255,))
        title = str(scene.get("text") or "").strip()
        if title:
            draw = ImageDraw.Draw(img)
            size_pt = max(28, w // 12)
            font = _font(size_pt)
            lines = _wrap(draw, title[:200], font, int(w * 0.85))
            line_h = int(size_pt * 1.25)
            y = (h - len(lines) * line_h) // 2
            for ln in lines:
                draw.text(((w - draw.textlength(ln, font=font)) / 2, y), ln, font=font, fill=(245, 243, 255))
                y += line_h
            return img
    caption = str(scene.get("caption") or scene.get("text") or "").strip()
    if caption:
        _caption(img, caption[:160], "bottom")
    return img


def _color(value: Any, default: tuple[int, int, int]) -> tuple[int, int, int]:
    m = re.fullmatch(r"#?([0-9a-fA-F]{6})", str(value or ""))
    if not m:
        return default
    v = m.group(1)
    return int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16)


async def video(scenes: Any, shape: str = "square", timeout: float = 180) -> MediaFile:
    """Ролик из сцен: картинка (SVG) или текст, плавное приближение и затухание. Без звука."""
    if not isinstance(scenes, list) or not scenes:
        raise MediaError("scenes: список сцен {svg? или text, caption?, seconds, bg?}")
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise MediaError("на сервере нет ffmpeg — видео пока делать не умею")
    scenes = [s for s in scenes if isinstance(s, dict)][:VIDEO_MAX_SCENES]
    size = VIDEO_SIZES.get(shape, VIDEO_SIZES["square"])
    secs = [min(max(float(s.get("seconds") or 3), 1.0), 8.0) for s in scenes]
    if sum(secs) > VIDEO_MAX_SECONDS:
        k = VIDEO_MAX_SECONDS / sum(secs)
        secs = [s * k for s in secs]
    try:
        cards = await asyncio.wait_for(asyncio.to_thread(lambda: [_scene_card(s, size) for s in scenes]), timeout)
    except asyncio.TimeoutError as exc:
        raise MediaError("рисунки для видео слишком сложные, сделай проще") from exc

    with tempfile.TemporaryDirectory(prefix="stella-video-") as tmp:
        out = os.path.join(tmp, "video.mp4")
        proc = await asyncio.create_subprocess_exec(
            ffmpeg, "-hide_banner", "-loglevel", "error", "-y",
            "-f", "rawvideo", "-pix_fmt", "rgb24", "-s", f"{size[0]}x{size[1]}", "-r", str(VIDEO_FPS), "-i", "-",
            "-c:v", "libx264", "-pix_fmt", "yuv420p", "-preset", "veryfast", "-crf", "26",
            "-movflags", "+faststart", out,
            stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.PIPE,
        )

        async def feed() -> None:
            assert proc.stdin is not None
            for card, seconds in zip(cards, secs, strict=True):
                base = card.convert("RGB")
                n = max(1, int(seconds * VIDEO_FPS))
                for start in range(0, n, VIDEO_FPS):  # по секунде за раз — не держим весь ролик в памяти
                    chunk = await asyncio.to_thread(_frames, base, start, min(start + VIDEO_FPS, n), n, size)
                    proc.stdin.write(chunk)
                    await proc.stdin.drain()
            proc.stdin.close()

        try:
            await asyncio.wait_for(feed(), timeout=timeout)
            _, err = await asyncio.wait_for(proc.communicate(), timeout=timeout)
        except (asyncio.TimeoutError, BrokenPipeError, ConnectionResetError) as exc:
            proc.kill()
            raise MediaError("видео не успело собраться, сделай короче") from exc
        if proc.returncode != 0:
            raise MediaError(f"ffmpeg: {(err or b'').decode(errors='replace')[:200]}")
        with open(out, "rb") as f:
            data = f.read()
    return MediaFile("video", data, MIME["mp4"], "video.mp4")


def _frames(base: Image.Image, start: int, stop: int, n: int, size: tuple[int, int]) -> bytes:
    """Кадры start..stop сцены из n кадров: медленное приближение (1.0 -> 1.08) и затухание по краям."""
    w, h = size
    fade = min(int(0.4 * VIDEO_FPS), n // 3)
    frames = []
    for i in range(start, stop):
        zoom = 1.0 + 0.08 * (i / max(n - 1, 1))
        cw, ch = int(w / zoom), int(h / zoom)
        x, y = (w - cw) // 2, (h - ch) // 2
        frame = base.crop((x, y, x + cw, y + ch)).resize((w, h), Image.BILINEAR)
        alpha = 1.0
        if fade and i < fade:
            alpha = (i + 1) / fade
        elif fade and i >= n - fade:
            alpha = (n - i) / fade
        if alpha < 1.0:
            frame = Image.eval(frame, lambda px, a=alpha: int(px * a))
        frames.append(frame.tobytes())
    return b"".join(frames)
