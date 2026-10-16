"""Стелла читает файлы и ссылки: все форматы, лимиты, zip-бомбы, мусор и защита от походов во внутреннюю сеть.

Интернет не нужен: DNS подменяется (loop.getaddrinfo), сайты — httpx.MockTransport. TLS с SNI проверяется
на своём локальном сервере. Живой тест с example.com — только с STELLA_LIVE_TESTS=1.
"""

import asyncio
import gzip
import io
import os
import random
import socket
import ssl
import tracemalloc
import zipfile
import zlib
from datetime import datetime, timedelta, timezone

import httpx
import pytest

from bot import reader
from bot.reader import ReaderError, extract_text, fetch_url

PUBLIC_IP = "93.184.215.14"
PUBLIC_IP2 = "93.184.215.15"


# ---------- файлы для тестов (делаем в памяти) ----------
def make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Доклад о космосе", level=1)
    doc.add_paragraph("до таблицы")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "планета", "спутник"
    table.cell(1, 0).text, table.cell(1, 1).text = "Земля", "Луна"
    merged = table.cell(2, 0).merge(table.cell(2, 1))
    merged.text = "объединённая ячейка"
    doc.add_paragraph("после таблицы")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Оценки"
    ws.append(["имя", "балл", "дата"])
    ws.append(["Вася", 5.0, datetime(2026, 9, 1)])
    ws.append(["Петя", 4.5, None])
    ws["D2"] = "=1+1"  # формула без сохранённого значения
    ws["E2"] = '=WEBSERVICE("http://169.254.169.254/")'
    ws.append([])
    ws["A7"] = "строка\tс табом"
    other = wb.create_sheet("Расходы")
    other.append(["хлеб", 50])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])  # заголовок и текст
    s.shapes.title.text = "Космос"
    s.placeholders[1].text = "Луна\nМарс"
    s.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)).text_frame.text = "надпись"
    group = s.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(5), Inches(5), Inches(2), Inches(1)).text_frame.text = "в группе"
    table = s.shapes.add_table(2, 2, Inches(5), Inches(1), Inches(3), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "звезда", "Солнце"
    s.notes_slide.notes_text_frame.text = "сказать про Луну"
    s2 = prs.slides.add_slide(prs.slide_layouts[0])
    s2.shapes.title.text = "Конец"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf(pages: int = 2, image: bool = False, user_password: str | None = None) -> bytes:
    from fpdf import FPDF

    from bot.media import font_path

    pdf = FPDF()
    font = font_path()
    if font:
        pdf.add_font("DejaVu", "", font)
        pdf.set_font("DejaVu", size=14)
    else:
        pdf.set_font("Helvetica", size=14)
    if user_password is not None:
        pdf.set_encryption(owner_password="owner-secret", user_password=user_password)
    for i in range(1, pages + 1):
        pdf.add_page()
        pdf.multi_cell(0, 10, f"Page {i} marker", new_x="LMARGIN", new_y="NEXT")
        if font:
            pdf.multi_cell(0, 10, f"Страница номер {i}", new_x="LMARGIN", new_y="NEXT")
        if image and i == 1:
            from PIL import Image

            buf = io.BytesIO()
            Image.new("RGB", (40, 40), (200, 30, 30)).save(buf, "PNG")
            buf.seek(0)
            pdf.image(buf, x=10, y=60, w=30)
    return bytes(pdf.output())


def docx_like_zip(document_xml_chunks, count: int = 1) -> bytes:
    """Zip с word/document.xml из кусков (пишем потоком, чтобы тест не ел память)."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        with zf.open("word/document.xml", "w") as f:
            for chunk in document_xml_chunks:
                f.write(chunk)
        for i in range(count - 2):
            zf.writestr(f"word/media/p{i}.xml", "")
    return out.getvalue()


def lie_about_size(data: bytes, member: str, fake_size: int) -> bytes:
    """Подделать оглавление zip: часть «весит» fake_size, и CRC подогнан под первые fake_size байт."""
    real = zipfile.ZipFile(io.BytesIO(data)).read(member)
    crc = zlib.crc32(real[:fake_size]).to_bytes(4, "little")
    size = fake_size.to_bytes(4, "little")
    buf = bytearray(data)
    name = member.encode()
    i = buf.find(b"PK\x01\x02")
    while i >= 0:
        name_len = int.from_bytes(buf[i + 28:i + 30], "little")
        if bytes(buf[i + 46:i + 46 + name_len]) == name:
            buf[i + 16:i + 20], buf[i + 24:i + 28] = crc, size  # центральное оглавление
            off = int.from_bytes(buf[i + 42:i + 46], "little")
            buf[off + 14:off + 18], buf[off + 22:off + 26] = crc, size  # локальный заголовок
        i = buf.find(b"PK\x01\x02", i + 4)
    return bytes(buf)


# ---------- форматы ----------
def test_docx_paragraphs_and_tables_in_order():
    text = extract_text(make_docx(), "доклад.docx",
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "Доклад о космосе" in text
    assert "планета | спутник" in text and "Земля | Луна" in text
    assert text.count("объединённая ячейка") == 1  # объединённые ячейки не дублируются
    assert text.index("до таблицы") < text.index("Земля | Луна") < text.index("после таблицы")


def test_xlsx_values_only_never_formulas():
    text = extract_text(make_xlsx(), "оценки.xlsx", None)
    assert "=== Лист «Оценки» ===" in text and "=== Лист «Расходы» ===" in text
    assert "имя\tбалл\tдата" in text
    assert "Вася\t5\t2026-09-01" in text and "Петя\t4.5" in text
    assert "строка с табом" in text  # таб внутри ячейки не ломает столбцы
    assert "хлеб\t50" in text
    assert "=1+1" not in text and "WEBSERVICE" not in text and "169.254" not in text
    assert "Вася\t5\t2026-09-01\t2" not in text  # формулу никто не посчитал


def test_pptx_titles_text_tables_groups_notes():
    text = extract_text(make_pptx(), "space.pptx", None)
    assert "--- слайд 1 ---" in text and "--- слайд 2 ---" in text
    for part in ("Луна", "Марс", "надпись", "в группе", "звезда | Солнце", "Заметки: сказать про Луну", "Конец"):
        assert part in text, part
    assert text.count("Космос") == 1  # заголовок не повторяется


def test_pdf_pages_and_images_skipped():
    text = extract_text(make_pdf(2, image=True), "doc.pdf", "application/pdf")
    assert "--- страница 1 ---" in text and "--- страница 2 ---" in text
    assert "Page 1 marker" in text and "Page 2 marker" in text
    from bot.media import font_path

    if font_path():
        assert "Страница номер 2" in text


def test_pdf_page_limit():
    text = extract_text(make_pdf(65), "big.pdf", None)
    assert "--- страница 60 ---" in text and "--- страница 61 ---" not in text
    assert "дальше обрезано: прочитала 60 страниц из 65" in text


def test_pdf_with_password():
    try:
        locked = make_pdf(1, user_password="secret")
        owner_only = make_pdf(1, user_password="")
    except Exception as exc:  # pragma: no cover - старый fpdf2
        pytest.skip(f"fpdf2 не шифрует: {exc}")
    with pytest.raises(ReaderError, match="под паролем"):
        extract_text(locked, "locked.pdf", None)
    assert "Page 1 marker" in extract_text(owner_only, "owner.pdf", None)


def test_pdf_found_by_content_even_with_wrong_name():
    assert "Page 1 marker" in extract_text(make_pdf(1), "file.bin", "application/octet-stream")


@pytest.mark.parametrize("data,name,mime", [
    ("Привет, мир! Ёлка".encode("cp1251"), "старый.txt", "text/plain"),
    ("Привет, мир! Ёлка".encode("cp1251"), "", "text/plain; charset=windows-1251"),
    ("Привет, мир! Ёлка".encode("utf-8"), "x.md", None),
    (b"\xef\xbb\xbf" + "Привет, мир! Ёлка".encode("utf-8"), "bom.txt", None),
    ("Привет, мир! Ёлка".encode("utf-16"), "u16.txt", None),
    ("имя;возраст\nПривет, мир! Ёлка;10\n".encode("cp1251"), "", "text/csv"),
])
def test_text_files_and_encodings(data, name, mime):
    assert "Привет, мир! Ёлка" in extract_text(data, name, mime)


def test_json_is_pretty_and_readable():
    text = extract_text(b'{"\\u0438\\u043c\\u044f": "\\u0412\\u0430\\u0441\\u044f", "n": [1, 2]}', "a.json", None)
    assert '"имя": "Вася"' in text
    assert extract_text(b"{broken json", "b.json", None) == "{broken json"


def test_html_file_visible_text_only():
    html = ('<html><head><meta charset="windows-1251"><title>Заголовок</title><style>p{color:red}</style>'
            '</head><body><nav>Меню сайта</nav><h1>Главное</h1><p>Видимый   текст</p>'
            '<div style="display: none">спрятано стилем</div><p hidden>спрятано атрибутом</p>'
            '<div aria-hidden="true"><div>и это</div>тоже</div><p>после скрытого</p>'
            '<script>var secret = "<p>не текст</p>";</script><table><tr><td>a</td><td>b</td></tr></table>'
            '<ul><li>раз</li><li>два</li></ul><footer>Подвал</footer></body></html>').encode("cp1251")
    text = extract_text(html, "page.html", "text/html")
    assert text.startswith("Заголовок\n")
    for part in ("Главное", "Видимый текст", "после скрытого", "a | b", "• раз\n• два"):
        assert part in text, part
    for bad in ("Меню", "спрятано", "и это", "тоже", "secret", "не текст", "Подвал", "color"):
        assert bad not in text, bad


# ---------- лимиты ----------
def test_output_is_cut_with_marker():
    text = extract_text(("слово " * 20_000).encode(), "long.txt", None)
    assert len(text) <= reader.MAX_OUTPUT_CHARS
    assert text.endswith("…(дальше обрезано)")


def test_input_size_limit():
    with pytest.raises(ReaderError, match="20 МБ"):
        extract_text(b"a" * (reader.MAX_INPUT_BYTES + 1), "huge.txt", None)


@pytest.mark.parametrize("data,name,mime,msg", [
    (b"\xff\xd8\xff\xe0" + bytes(200), "photo.jpg", "image/jpeg", "пока не читаю"),
    (b"\x1a\x45\xdf\xa3" + bytes(200), "video", "video/webm", "пока не читаю"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(600), "old.doc", "application/msword", "старый формат"),
    (b"{\\rtf1 hello}", "x.rtf", "application/rtf", "старый формат"),
])
def test_unknown_types(data, name, mime, msg):
    with pytest.raises(ReaderError, match=msg):
        extract_text(data, name, mime)


def test_plain_zip_archive_is_not_read():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", "привет")
    with pytest.raises(ReaderError, match="пока не читаю"):
        extract_text(buf.getvalue(), "archive.zip", "application/zip")


# ---------- zip-бомбы ----------
def test_zip_bomb_by_declared_size():
    mb = b" " * (1 << 20)
    bomb = docx_like_zip(mb for _ in range(101))  # 101 МБ пробелов сжимаются в ~100 КБ
    assert len(bomb) < 1_000_000
    with pytest.raises(ReaderError, match="zip-бомб"):
        extract_text(bomb, "bomb.docx", None)


def test_zip_bomb_by_member_count():
    bomb = docx_like_zip([b"<w:document/>"], count=reader.MAX_ZIP_MEMBERS + 10)
    with pytest.raises(ReaderError, match="много частей"):
        extract_text(bomb, "many.xlsx", None)


def test_zip_bomb_with_forged_sizes_does_not_eat_memory():
    """Оглавление врёт (часть «весит» 1000 байт, CRC подогнан), а внутри 30 МБ.

    Обычный zipfile.read() распаковал бы всё целиком — мы ловим это до разбора, почти без памяти."""
    real = docx_like_zip(b"\0" * (1 << 20) for _ in range(30))
    forged = lie_about_size(real, "word/document.xml", 1000)
    assert len(zipfile.ZipFile(io.BytesIO(forged)).read("word/document.xml")) == 1000  # zipfile поверил
    tracemalloc.start()
    try:
        with pytest.raises(ReaderError, match="zip-бомба|повреждён"):
            extract_text(forged, "forged.docx", None)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert peak < 8 * 1024 * 1024


# ---------- мусор ----------
def _garbage_cases():
    rnd = random.Random(42)
    junk = rnd.randbytes(5000)
    docx, xlsx, pptx = make_docx(), make_xlsx(), make_pptx()
    cases = [(junk, n, m) for n, m in [
        ("x.pdf", "application/pdf"), ("x.docx", None), ("x.xlsx", None), ("x.pptx", None), ("x.txt", "text/plain"),
        ("x.json", None), ("x.html", "text/html"), ("x.csv", None), ("", None), ("x.bin", None),
    ]]
    cases += [
        (b"%PDF-1.7\n" + junk, "x.pdf", None),
        (b"%PDF-1.4\n1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\ntrailer << /Root 1 0 R >>\n%%EOF",
         "x.pdf", None),
        (b"PK\x03\x04" + junk, "x.docx", None),
        (b"PK\x05\x06" + bytes(18), "x.xlsx", None),
        (docx[: len(docx) // 2], "half.docx", None),
        (xlsx[: len(xlsx) // 2], "half.xlsx", None),
        (pptx[: len(pptx) // 2], "half.pptx", None),
        (docx, "x.xlsx", None),  # не то расширение: определим по содержимому — и прочитаем
        (bytes(1000), "zeros.txt", None),
        (b"", "empty.txt", None),
        (b"   \n\n  ", "blank.txt", None),
    ]
    return cases


@pytest.mark.parametrize("data,name,mime", _garbage_cases())
def test_garbage_never_crashes(data, name, mime):
    try:
        text = extract_text(data, name, mime)
    except ReaderError as exc:
        assert str(exc) and len(str(exc)) < 200
    else:
        assert isinstance(text, str) and len(text) <= reader.MAX_OUTPUT_CHARS


# ---------- ссылки: подменённые DNS и сайты ----------
class FakeWeb:
    def __init__(self) -> None:
        self.dns: dict[str, list[str]] = {}
        self.pages: dict[str, object] = {}
        self.requests: list[httpx.Request] = []
        self.lookups: list[str] = []

    async def getaddrinfo(self, host, port, *, family=0, type=0, proto=0, flags=0):
        self.lookups.append(host)
        ips = self.dns.get(host, [PUBLIC_IP])
        if not ips:
            raise socket.gaierror(socket.EAI_NONAME, "Name or service not known")
        return [(socket.AF_INET6, socket.SOCK_STREAM, 6, "", (ip, port, 0, 0)) if ":" in ip
                else (socket.AF_INET, socket.SOCK_STREAM, 6, "", (ip, port)) for ip in ips]

    async def handle(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        key = f"{request.url.scheme}://{request.headers['host']}{request.url.raw_path.decode()}"
        make = self.pages.get(key)
        if make is None:
            return httpx.Response(404, headers={"content-type": "text/html"}, content=b"<p>nope</p>")
        result = make(request)
        return await result if asyncio.iscoroutine(result) else result


def page(body: bytes | str, ctype: str = "text/html; charset=utf-8", status: int = 200, **headers):
    data = body.encode() if isinstance(body, str) else body
    # stream=, а не content=: так тело читается кусками, как у настоящего ответа из сети
    return lambda req: httpx.Response(status, headers={"content-type": ctype, "content-length": str(len(data)),
                                                       **headers}, stream=httpx.ByteStream(data))


def redirect(location: str, status: int = 302):
    return lambda req: httpx.Response(status, headers={"location": location})


@pytest.fixture
async def net(monkeypatch):
    web = FakeWeb()
    monkeypatch.setattr(asyncio.get_running_loop(), "getaddrinfo", web.getaddrinfo)
    monkeypatch.setattr(reader, "_transport", lambda: httpx.MockTransport(web.handle))
    return web


ARTICLE = ("<html><head><title>Новости</title><script>steal()</script></head><body><nav>Главная | О нас</nav>"
           "<article><h1>Стелла научилась читать</h1><p>Теперь она открывает ссылки.</p></article>"
           "<footer>© 2026</footer></body></html>")


async def test_fetch_pins_vetted_ip_with_host_and_sni(net):
    net.pages["https://example.com/news?id=1"] = page(ARTICLE)
    final, text = await fetch_url("https://example.com/news?id=1#top")
    assert final == "https://example.com/news?id=1"
    assert text.startswith("Новости\n") and "Стелла научилась читать" in text and "открывает ссылки" in text
    for bad in ("steal", "Главная", "©"):
        assert bad not in text
    (req,) = net.requests
    assert req.url.host == PUBLIC_IP  # подключаемся к проверенному IP, а не к имени
    assert req.headers["host"] == "example.com"
    assert req.extensions["sni_hostname"] == "example.com"  # сертификат проверится по настоящему имени
    assert req.headers["user-agent"] == "Mozilla/5.0 (compatible; StellaBot/1.0)"
    assert net.lookups == ["example.com"]


async def test_fetch_http_with_port_and_no_sni(net):
    net.pages["http://example.com:8080/x"] = page("просто текст", "text/plain")
    final, text = await fetch_url("http://example.com:8080/x")
    assert (final, text) == ("http://example.com:8080/x", "просто текст")
    (req,) = net.requests
    assert req.url.host == PUBLIC_IP and req.url.port == 8080
    assert req.headers["host"] == "example.com:8080"
    assert "sni_hostname" not in req.extensions


async def test_fetch_without_scheme_and_json(net):
    net.pages["https://api.example.com/data"] = page('{"\\u0438\\u043c\\u044f": "Вася"}', "application/json")
    final, text = await fetch_url("api.example.com/data")
    assert final == "https://api.example.com/data" and '"имя": "Вася"' in text


def test_pin_ipv6():
    url, headers, ext = reader._pin(httpx.URL("https://example.com/a?b=1"), "2606:2800:21f::1")
    assert str(url) == "https://[2606:2800:21f::1]/a?b=1"
    assert headers == {"Host": "example.com"} and ext == {"sni_hostname": "example.com"}


BLOCKED = [
    "http://localhost/", "http://LOCALHOST:8080/", "http://localhost./", "http://127.0.0.1/", "http://127.0.0.1:8000/",
    "http://10.0.0.5/", "http://10.255.255.255:8080/", "http://172.16.3.4/", "http://172.31.0.1/",
    "http://192.168.1.1/", "http://169.254.169.254/latest/meta-data/", "http://[::1]/", "http://[::1]:8080/",
    "http://[::ffff:127.0.0.1]/", "http://[::ffff:7f00:1]/", "http://[::ffff:8.8.8.8]/", "http://[fe80::1%25eth0]/",
    "http://[fd00::1]/", "http://[2002:7f00:1::1]/", "http://0.0.0.0/", "http://0.0.0.0:8000/", "http://[::]/",
    "http://2130706433/", "http://0x7f.1/", "http://0x7f000001/", "http://127.1/", "http://0177.0.0.1/",
    "http://100.64.0.1/", "http://224.0.0.1/", "http://255.255.255.255/", "http://198.18.0.1/",
    "http://evolution:8080/instance/fetchInstances", "http://postgres/", "http://redis/", "http://bot:8000/health",
    "http://stella-bot/", "http://host.docker.internal/", "http://metadata.google.internal/", "http://router.lan/",
    "http://printer.local/", "http://something.localhost/",
    "https://example.com:22/", "http://example.com:6379/", "http://example.com:5432/", "http://example.com:8000/",
    "file:///etc/passwd", "ftp://example.com/x", "gopher://example.com/", "javascript:alert(1)",
    "data:text/html,<p>hi</p>", "http://user:pass@example.com/", "http://user@example.com/",
    "", "http://", "http://" + "a" * 3000 + ".com/",
]


@pytest.mark.parametrize("url", BLOCKED)
async def test_internal_and_weird_urls_are_blocked(net, url):
    with pytest.raises(ReaderError, match="^не получилось открыть ссылку: "):
        await fetch_url(url)
    assert net.requests == []


@pytest.mark.parametrize("ips", [
    ["10.0.0.7"], ["172.20.0.3"], ["192.168.0.10"], ["127.0.0.1"], ["169.254.169.254"], ["100.100.100.200"],
    [PUBLIC_IP, "127.0.0.1"],  # хотя бы один внутренний — не идём совсем
    ["::1"], ["::ffff:192.168.0.1"], ["fd12:3456::1"], ["fe80::1"], ["0.0.0.0"],
])
async def test_hostname_resolving_to_internal_ip_is_blocked(net, ips):
    net.dns["innocent.example.org"] = ips
    with pytest.raises(ReaderError, match="внутреннюю сеть"):
        await fetch_url("https://innocent.example.org/")
    assert net.requests == []


async def test_unknown_host(net):
    net.dns["nope.example.org"] = []
    with pytest.raises(ReaderError, match="такого сайта нет"):
        await fetch_url("https://nope.example.org/")


@pytest.mark.parametrize("location", [
    "http://127.0.0.1:8080/admin", "http://[::1]/", "http://169.254.169.254/latest/meta-data/",
    "http://evolution:8080/", "http://localhost/", "http://2130706433/", "https://sneaky.example.net/",
    "http://example.com:6379/",
])
async def test_redirect_to_internal_is_blocked(net, location):
    net.dns["sneaky.example.net"] = ["10.1.2.3"]
    net.pages["https://example.com/go"] = redirect(location)
    with pytest.raises(ReaderError, match="^не получилось открыть ссылку: "):
        await fetch_url("https://example.com/go")
    assert len(net.requests) == 1  # во внутреннюю сеть не сходили


@pytest.mark.parametrize("location", ["file:///etc/passwd", "ftp://example.com/x", "gopher://example.com/",
                                      "javascript:alert(1)", "http://example.com/plain"])
async def test_redirect_to_other_scheme_is_blocked(net, location):
    net.pages["https://example.com/go"] = redirect(location, 301)
    with pytest.raises(ReaderError, match="^не получилось открыть ссылку: сайт переадресует"):
        await fetch_url("https://example.com/go")
    assert len(net.requests) == 1


async def test_redirects_are_followed_and_rechecked(net):
    net.dns["www.example.com"] = [PUBLIC_IP2]
    net.pages["http://example.com/"] = redirect("https://www.example.com/start", 301)  # http -> https можно
    net.pages["https://www.example.com/start"] = redirect("/final?x=1", 307)
    net.pages["https://www.example.com/final?x=1"] = page("<title>Финиш</title><p>дошли</p>")
    final, text = await fetch_url("http://example.com/")
    assert final == "https://www.example.com/final?x=1"
    assert "дошли" in text
    assert [r.url.host for r in net.requests] == [PUBLIC_IP, PUBLIC_IP2, PUBLIC_IP2]
    assert net.requests[1].headers["host"] == "www.example.com"
    assert net.requests[1].extensions["sni_hostname"] == "www.example.com"
    assert net.lookups == ["example.com", "www.example.com", "www.example.com"]  # каждый шаг — заново


async def test_too_many_redirects(net):
    for i in range(6):
        net.pages[f"https://example.com/r{i}"] = redirect(f"/r{i + 1}")
    with pytest.raises(ReaderError, match="слишком много переадресаций"):
        await fetch_url("https://example.com/r0")
    assert len(net.requests) == reader.MAX_REDIRECTS + 1


async def test_body_is_capped_while_streaming(net):
    pulled = 0

    async def endless():
        nonlocal pulled
        while True:
            pulled += 1
            yield ("бесконечный текст " * 1000).encode()

    net.pages["https://example.com/endless"] = lambda req: httpx.Response(
        200, headers={"content-type": "text/plain; charset=utf-8"}, content=endless())
    _, text = await fetch_url("https://example.com/endless", max_bytes=200_000)
    assert len(text) <= reader.MAX_PAGE_CHARS and text.endswith("…(дальше обрезано)")
    assert pulled * len(("бесконечный текст " * 1000).encode()) < 200_000 + 70_000


async def test_gzip_bomb_from_site_is_capped(net):
    bomb = gzip.compress(b"a" * 50_000_000)
    net.pages["https://example.com/bomb"] = page(bomb, "text/plain", **{"content-encoding": "gzip"})
    tracemalloc.start()
    try:
        _, text = await fetch_url("https://example.com/bomb", max_bytes=1_000_000)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert text.startswith("aaaa") and len(text) <= reader.MAX_PAGE_CHARS
    assert peak < 20_000_000


async def test_gzip_page_is_unpacked(net):
    net.pages["https://example.com/gz"] = page(gzip.compress("<p>сжатая страница</p>".encode()),
                                               **{"content-encoding": "gzip"})
    _, text = await fetch_url("https://example.com/gz")
    assert text == "сжатая страница"


async def test_pdf_by_link(net):
    net.pages["https://example.com/doc.pdf"] = page(make_pdf(2), "application/pdf")
    _, text = await fetch_url("https://example.com/doc.pdf")
    assert "Page 2 marker" in text


async def test_pdf_by_link_too_big(net):
    pdf = make_pdf(1)
    net.pages["https://example.com/big.pdf"] = page(pdf, "application/pdf")

    async def stream():
        yield pdf
        yield b"0" * 50_000

    net.pages["https://example.com/stream.pdf"] = lambda req: httpx.Response(
        200, headers={"content-type": "application/pdf"}, content=stream())
    for url in ("https://example.com/big.pdf", "https://example.com/stream.pdf"):
        with pytest.raises(ReaderError, match="PDF слишком большой"):
            await fetch_url(url, max_bytes=len(pdf) - 10)


@pytest.mark.parametrize("ctype", ["image/png", "application/octet-stream", "application/zip", "video/mp4",
                                   "application/javascript", "text/css", None])
async def test_content_type_is_rejected(net, ctype):
    headers = {"content-type": ctype} if ctype else {}
    net.pages["https://example.com/file"] = lambda req: httpx.Response(200, headers=headers, content=b"\x89PNG....")
    with pytest.raises(ReaderError, match="там не страница|не сказал"):
        await fetch_url("https://example.com/file")


async def test_http_error_status(net):
    with pytest.raises(ReaderError, match="сайт ответил 404"):
        await fetch_url("https://example.com/missing")


async def test_timeout(net):
    async def slow(req):
        await asyncio.sleep(5)
        return httpx.Response(200, headers={"content-type": "text/plain"}, content=b"late")

    net.pages["https://example.com/slow"] = slow
    with pytest.raises(ReaderError, match="долго"):
        await fetch_url("https://example.com/slow", timeout=0.3)


async def test_connect_error_tries_next_ip(net):
    net.dns["down.example.com"] = [PUBLIC_IP, PUBLIC_IP2]

    def down(req):
        raise httpx.ConnectError("connection refused", request=req)

    net.pages["https://down.example.com/"] = down
    with pytest.raises(ReaderError, match="сайт недоступен"):
        await fetch_url("https://down.example.com/")
    assert [r.url.host for r in net.requests] == [PUBLIC_IP, PUBLIC_IP2]


async def test_page_text_is_cut(net):
    net.pages["https://example.com/long"] = page("<p>" + "очень длинный текст " * 5000 + "</p>")
    _, text = await fetch_url("https://example.com/long")
    assert len(text) <= reader.MAX_PAGE_CHARS and text.endswith("…(дальше обрезано)")


# ---------- инструмент для модели ----------
async def test_read_link_tool(net):
    from bot.tools import ToolError

    net.pages["https://example.com/"] = page(ARTICLE)
    (tool,) = reader.build_reader_tools(None, None, None)
    assert tool.name == "read_link" and tool.input_schema["required"] == ["url"]
    out = await tool.handler({"url": "https://example.com/"})
    assert out.startswith("Страница https://example.com/:\nНовости") and "Стелла научилась читать" in out
    with pytest.raises(ToolError, match="^не получилось открыть ссылку: это внутренний адрес"):
        await tool.handler({"url": "http://127.0.0.1/"})
    with pytest.raises(ToolError, match="нужна ссылка"):
        await tool.handler({})
    await tool.handler({"url": "https://example.com/"})
    with pytest.raises(ToolError, match="хватит трёх"):
        await tool.handler({"url": "https://example.com/"})


# ---------- настоящий TLS: подключение к IP, а сертификат проверяется по имени (SNI) ----------
TLS_HOST = "reader.stella-bot.net"


def _self_signed(tmp_path, hostname: str):
    x509 = pytest.importorskip("cryptography.x509")
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import ec
    from cryptography.x509.oid import ExtendedKeyUsageOID, NameOID

    key = ec.generate_private_key(ec.SECP256R1())
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, hostname)])
    now = datetime.now(timezone.utc)
    ski = x509.SubjectKeyIdentifier.from_public_key(key.public_key())
    cert = (
        x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key())
        .serial_number(x509.random_serial_number())
        .not_valid_before(now - timedelta(days=1)).not_valid_after(now + timedelta(days=1))
        .add_extension(x509.SubjectAlternativeName([x509.DNSName(hostname)]), critical=False)
        .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
        .add_extension(x509.KeyUsage(True, False, False, False, False, True, False, False, False), critical=True)
        .add_extension(x509.ExtendedKeyUsage([ExtendedKeyUsageOID.SERVER_AUTH]), critical=False)
        .add_extension(ski, critical=False)
        .add_extension(x509.AuthorityKeyIdentifier.from_issuer_subject_key_identifier(ski), critical=False)
        .sign(key, hashes.SHA256())
    )
    cert_file, key_file = tmp_path / "cert.pem", tmp_path / "key.pem"
    cert_file.write_bytes(cert.public_bytes(serialization.Encoding.PEM))
    key_file.write_bytes(key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                           serialization.NoEncryption()))
    return str(cert_file), str(key_file)


@pytest.fixture
async def tls_site(tmp_path, monkeypatch):
    """Локальный HTTPS-сайт с сертификатом на TLS_HOST. Запоминает SNI и Host каждого запроса."""
    cert_file, key_file = _self_signed(tmp_path, TLS_HOST)
    server_ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    server_ctx.load_cert_chain(cert_file, key_file)
    seen: dict[str, list] = {"sni": [], "host": []}
    server_ctx.sni_callback = lambda sslobj, name, ctx: seen["sni"].append(name)

    async def handle(r: asyncio.StreamReader, w: asyncio.StreamWriter) -> None:
        try:
            head = await r.readuntil(b"\r\n\r\n")
            for line in head.decode().split("\r\n"):
                if line.lower().startswith("host:"):
                    seen["host"].append(line.split(":", 1)[1].strip())
            body = "<title>TLS</title><p>секретов нет</p>".encode()
            w.write(b"HTTP/1.1 200 OK\r\nContent-Type: text/html; charset=utf-8\r\nConnection: close\r\n"
                    b"Content-Length: " + str(len(body)).encode() + b"\r\n\r\n" + body)
            await w.drain()
        except (ConnectionError, asyncio.IncompleteReadError, ssl.SSLError):
            pass
        finally:
            w.close()

    server = await asyncio.start_server(handle, "127.0.0.1", 0, ssl=server_ctx)
    port = server.sockets[0].getsockname()[1]
    client_ctx = ssl.create_default_context(cafile=cert_file)

    async def resolve(url):  # «сайт» живёт на 127.0.0.1 — обычную проверку адреса тут обходим
        return ["127.0.0.1"]

    monkeypatch.setattr(reader, "_resolve", resolve)
    monkeypatch.setattr(reader, "ALLOWED_PORTS", frozenset({port}))
    monkeypatch.setattr(reader, "_transport", lambda: httpx.AsyncHTTPTransport(verify=client_ctx, trust_env=False))
    yield port, seen
    server.close()
    await server.wait_closed()


async def test_tls_pinned_ip_verifies_real_hostname(tls_site):
    port, seen = tls_site
    final, text = await fetch_url(f"https://{TLS_HOST}:{port}/")
    assert final == f"https://{TLS_HOST}:{port}/" and "секретов нет" in text
    assert seen["sni"] == [TLS_HOST] and seen["host"] == [f"{TLS_HOST}:{port}"]


async def test_tls_wrong_hostname_fails_verification(tls_site):
    port, seen = tls_site
    with pytest.raises(ReaderError, match="сертификат"):  # тот же IP, но сертификат не на это имя
        await fetch_url(f"https://other-name.stella-bot.net:{port}/")
    assert seen["sni"] == ["other-name.stella-bot.net"]


# ---------- живой тест (по желанию) ----------
@pytest.mark.skipif(not os.environ.get("STELLA_LIVE_TESTS"), reason="живой тест: STELLA_LIVE_TESTS=1")
async def test_live_example_com(monkeypatch):
    ca = os.environ.get("STELLA_LIVE_CA")  # если исходящий TLS идёт через прокси со своим сертификатом
    if ca:
        ctx = ssl.create_default_context(cafile=ca)
        monkeypatch.setattr(reader, "_transport", lambda: httpx.AsyncHTTPTransport(verify=ctx, trust_env=False))
    try:
        final, text = await fetch_url("https://example.com/")
    except ReaderError as exc:
        if "недоступен" in str(exc) or "долго" in str(exc) or "такого сайта нет" in str(exc):
            pytest.skip(f"нет интернета: {exc}")
        raise
    assert final == "https://example.com/" and "Example Domain" in text
