"""HTML-страницы: студия пишет страницу по заданию Стеллы, библиотеки и шрифты встраиваются в файл."""

import asyncio
import hashlib
import io
from datetime import date
from types import SimpleNamespace

import pytest
from PIL import Image

from bot import html_studio, media, weblibs
from bot.tools import ToolError

from .conftest import OWNER, settle, upsert

GROUP = "120363000000000001@g.us"
PAGE = "<!doctype html><html><head><title>Игра</title></head><body><h1>Привет</h1></body></html>"
BRIEF = "игра «поймай звезду» для чата, неоновый космос, счёт и рекорд, кнопка Играть"


def _asset(tmp_path, file, path, data):
    (tmp_path / file).write_bytes(data)
    return weblibs.Asset(file, path, hashlib.sha256(data).hexdigest())


# ---------- библиотеки и шрифты ----------
def test_resolve_names_and_cdn_addresses():
    assert weblibs.resolve("lib/three.js") == (weblibs.THREE,)
    assert weblibs.resolve("./lib/three.min.js") == (weblibs.THREE,)
    assert weblibs.resolve("lib/three-bloom.js")[-1] == weblibs.BLOOM_PASS
    assert weblibs.resolve("lib/OrbitControls.js") == (weblibs.ORBIT,)
    assert weblibs.resolve("https://cdnjs.cloudflare.com/ajax/libs/three.js/r128/three.min.js") == (weblibs.THREE,)
    assert weblibs.resolve("https://cdn.jsdelivr.net/npm/three@0.128.0/examples/js/controls/OrbitControls.js") == (
        weblibs.ORBIT,)
    assert weblibs.resolve("https://cdnjs.cloudflare.com/ajax/libs/gsap/3.12.2/ScrollTrigger.min.js") == (
        weblibs.SCROLL_TRIGGER,)
    assert weblibs.resolve("https://cdn.jsdelivr.net/npm/canvas-confetti@1.6.0/dist/confetti.browser.min.js") == (
        weblibs.CONFETTI,)
    assert weblibs.resolve("https://unpkg.com/matter-js@0.19.0/build/matter.min.js") == (weblibs.MATTER,)
    for other in ("https://evil.example/three.min.js", "lib/nope.js", "https://cdn.jsdelivr.net/npm/x@1/gsap.js",
                  "data:text/javascript,alert(1)"):
        assert weblibs.resolve(other) is None, other


def test_every_pinned_asset_is_a_jsdelivr_file_with_sha256():
    assert len(weblibs.ALL) == len({a.file for a in weblibs.ALL})  # имена в кеше не пересекаются
    for a in weblibs.ALL:
        assert len(a.sha256) == 64 and a.url.startswith("https://cdn.jsdelivr.net/npm/") and "@" in a.path
        assert a.sri.startswith("sha256-")


def test_libraries_are_embedded_or_loaded_with_integrity(tmp_path, monkeypatch):
    gsap = _asset(tmp_path, "gsap.min.js", "gsap@9/dist/gsap.min.js", b"window.gsap={v:'</script>'};")
    monkeypatch.setitem(weblibs.LIBS, "gsap", (gsap,))
    html = ('<script src="lib/gsap.js"></script><script src="lib/gsap.js"></script>'
            '<script src="lib/three.js"></script><script src="https://evil.example/x.js"></script><h1>hi</h1>')
    text = media.html_page("t", html, libs_dir=str(tmp_path)).data.decode()
    assert text.count("window.gsap=") == 1 and "<\\/script>'" in text  # один раз и не рвёт <script>
    assert f'src="{weblibs.THREE.url}" integrity="{weblibs.THREE.sri}"' in text  # нет в кеше — с CDN по хешу
    csp = text.split('content="', 1)[1].split('"', 1)[0]
    assert weblibs.THREE.url in csp and "evil.example" not in csp and "connect-src 'none'" in csp
    assert text.index("Content-Security-Policy") < text.index("window.gsap")


def test_fonts_only_when_used(tmp_path, monkeypatch):
    cyr = _asset(tmp_path, "inter-cyrillic.woff2", "x@1/inter-c.woff2", b"wOF2cyr")
    lat = weblibs.Asset("inter-latin.woff2", "x@1/inter-l.woff2", "0" * 64)  # нет в кеше
    monkeypatch.setitem(weblibs.FONTS, "Inter", ("100 900", (cyr, lat)))
    used = media.html_page("t", "<style>body{font-family:'Inter',sans-serif}</style><p>Interactive</p>",
                           libs_dir=str(tmp_path)).data.decode()
    assert "font-family:'Inter'" in used and "data:font/woff2;base64," in used and lat.url in used
    assert lat.url in used.split('content="', 1)[1].split('"', 1)[0]  # шрифт с CDN разрешён точечно
    plain = media.html_page("t", "<p>Interactive international</p>", libs_dir=str(tmp_path)).data.decode()
    assert "@font-face" not in plain
    gfonts = media.html_page("t", '<link href="https://fonts.googleapis.com/css2?family=Inter" rel="stylesheet">'
                                  "<style>@import url('https://fonts.googleapis.com/css2?family=Inter');</style>"
                                  "<p style=\"font-family: Inter\">x</p>", libs_dir=str(tmp_path)).data.decode()
    assert "googleapis" not in gfonts and "@font-face" in gfonts


async def test_ensure_downloads_checks_hash_and_caches(tmp_path, monkeypatch):
    import httpx

    good, bad = b"good-lib", b"tampered"
    a1 = weblibs.Asset("a.js", "a@1/a.js", hashlib.sha256(good).hexdigest())
    a2 = weblibs.Asset("b.js", "b@1/b.js", hashlib.sha256(b"expected").hexdigest())
    monkeypatch.setattr(weblibs, "ALL", (a1, a2))
    hits = []

    def handler(request):
        hits.append(request.url.path)
        return httpx.Response(200, content=good if request.url.path.endswith("a.js") else bad)

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        failed = await weblibs.ensure(str(tmp_path), client=client)
        assert failed == [a2] and (tmp_path / "a.js").read_bytes() == good and not (tmp_path / "b.js").exists()
        await weblibs.ensure(str(tmp_path), client=client)
    assert hits.count("/npm/a@1/a.js") == 1  # скачанное второй раз не качается


def test_photo_is_put_into_the_page():
    buf = io.BytesIO()
    Image.new("RGB", (3000, 2000), (200, 30, 30)).save(buf, "JPEG")
    photo = media.photo_for_page(buf.getvalue())
    assert photo[0] == "image/jpeg" and max(Image.open(io.BytesIO(photo[1])).size) == media.PAGE_PHOTO_SIDE
    text = media.html_page("t", '<img src="__PHOTO__"><div style="background:url(__PHOTO__)"></div>', photo).data
    assert text.count(b"data:image/jpeg;base64,") == 2 and b"__PHOTO__" not in text
    assert media.photo_for_page(b"not an image") is None
    sticker = io.BytesIO()
    Image.new("RGBA", (512, 512), (0, 0, 0, 0)).save(sticker, "WEBP")
    assert media.photo_for_page(sticker.getvalue())[0] == "image/webp"  # прозрачность сохраняется


def test_page_is_still_a_closed_box():
    text = media.html_page("Игра <3", '<meta http-equiv="refresh" content="0;url=http://evil"><base href="http://e/">'
                                      "<script>x=1</script>").data.decode()
    assert text.startswith("<!doctype html>") and text.index("Content-Security-Policy") < text.index("<script>")
    assert "refresh" not in text and "<base" not in text and "<title>Игра &lt;3</title>" in text
    for rule in ("default-src 'none'", "connect-src 'none'", "form-action 'none'", "worker-src blob:"):
        assert rule in text


# ---------- ответ студии ----------
def test_extract_html_variants():
    assert html_studio.extract_html(PAGE) == PAGE
    assert html_studio.extract_html(f"Вот страница:\n```html\n{PAGE}\n```\nГотово!") == PAGE
    assert html_studio.extract_html(f"Конечно! {PAGE} надеюсь понравится") == PAGE
    assert html_studio.extract_html("<style>b{}</style><div>x</div>").startswith("<style>")
    with pytest.raises(html_studio.StudioError):
        html_studio.extract_html("извини, не могу")


def test_eta_words():
    assert html_studio.human_eta(50) == "примерно через минуту"
    assert html_studio.human_eta(180) == "минуты через 3"
    assert html_studio.human_eta(400) == "минут через 7"


class FakeStream:
    def __init__(self, owner, kwargs):
        self.owner, self.kwargs = owner, kwargs

    async def __aenter__(self):
        self.owner.calls.append(self.kwargs)
        return self

    async def __aexit__(self, *exc):
        return False

    async def get_final_message(self):
        await asyncio.sleep(0)
        return self.owner.reply


class FakeClient:
    """replies — по ответу (текст, stop_reason) на каждый запрос; последний повторяется."""

    def __init__(self, text=PAGE, stop="end_turn", replies=None):
        self.calls = []
        self.replies = replies or [(text, stop)]
        self.messages = SimpleNamespace(stream=lambda **kw: FakeStream(self, kw))
        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=lambda **kw: FakeStream(self, kw)))

    @property
    def reply(self):
        text, stop = self.replies[min(len(self.calls), len(self.replies)) - 1]
        return SimpleNamespace(stop_reason=stop, content=[SimpleNamespace(type="thinking", thinking=""),
                                                          SimpleNamespace(type="text", text=text)],
                               usage=SimpleNamespace(input_tokens=3000, output_tokens=20000))


async def test_generate_asks_the_strong_model_with_design_prompt(settings, monkeypatch):
    fake = FakeClient(f"```html\n{PAGE}\n```")
    monkeypatch.setattr(html_studio, "_client", fake)
    res = await html_studio.generate(settings, BRIEF, "Звёзды", ("image/jpeg", b"\xff\xd8jpeg"))
    call = fake.calls[0]
    assert res.html == PAGE and res.cost_usd > 0.3
    assert call["model"] == "claude-opus-5-5" and call["thinking"] == {"type": "adaptive"}
    assert call["output_config"] == {"effort": "max", "task_budget": {"type": "tokens", "total": 100000}}
    assert call["betas"] == ["task-budgets-2026-03-13"] and call["max_tokens"] == 128000
    assert "lib/three.js" in call["system"] and "Unbounded" in call["system"]
    image, text = call["messages"][0]["content"]
    assert image["type"] == "image" and "__PHOTO__" in text["text"] and BRIEF in text["text"]


async def test_generate_errors_keep_the_cost(settings, monkeypatch):
    cut = PAGE.replace("</html>", "")  # оборвалось на полуслове
    fake = FakeClient(replies=[(cut, "max_tokens")])
    monkeypatch.setattr(html_studio, "_client", fake)
    with pytest.raises(html_studio.StudioError, match="слишком большой") as exc:
        await html_studio.generate(settings, BRIEF, "x")
    assert exc.value.cost_usd > 0.8  # обе попытки посчитаны
    assert [c["output_config"]["effort"] for c in fake.calls] == ["max", "high"]
    monkeypatch.setattr(html_studio, "_client", FakeClient("не буду", stop="refusal"))
    with pytest.raises(html_studio.StudioError, match="отказались"):
        await html_studio.generate(settings, BRIEF, "x")


async def test_polish_call_sees_draft_screenshots_and_errors(settings, monkeypatch):
    from bot.page_shots import Shots

    fake = FakeClient(PAGE)
    monkeypatch.setattr(html_studio, "_client", fake)
    shots = Shots([("сразу после открытия (1.5 с)", b"jpeg1"), ("после прокрутки", b"jpeg2")],
                  ["ошибка JS: TypeError: x is null"], overflow=True)
    await html_studio.generate(settings, BRIEF, "x", draft="<h1>черновик</h1>", shots=shots)
    call = fake.calls[0]
    assert call["output_config"]["effort"] == settings.html_polish_effort == "max"  # доводка — со своим усилием
    blocks = call["messages"][0]["content"]
    texts = "\n".join(b["text"] for b in blocks if b["type"] == "text")
    assert "<h1>черновик</h1>" in texts and "TypeError: x is null" in texts and "горизонтальная прокрутка" in texts
    assert "арт-директора" in texts and "Скриншот — после прокрутки" in texts
    assert [b["source"]["data"] for b in blocks if b["type"] == "image"] == ["anBlZzE=", "anBlZzI="]


@pytest.fixture
def polish_env(monkeypatch):
    from bot import page_shots

    calls = {"gen": [], "shots": [], "after_errors": 0, "polish_fail": False, "shots_fail": False}

    async def gen(settings, brief, title, image=None, *, draft=None, shots=None, kind="page"):
        calls["gen"].append(draft)
        if draft is not None and calls["polish_fail"]:
            raise html_studio.StudioError("слишком большой", 0.7)
        return html_studio.StudioResult("<h1>v2</h1>" if draft else "<h1>v1</h1>", 1.0 if draft else 2.0, 1.0)

    async def shoot(page, timeout=0):
        calls["shots"].append(page)
        if calls["shots_fail"]:
            raise page_shots.ShotsError("браузер упал")
        n = 0 if len(calls["shots"]) == 1 else calls["after_errors"]
        return page_shots.Shots([("кадр", b"j")], ["ошибка JS: boom"] * n)

    async def no_download(*a, **kw):
        return []

    monkeypatch.setattr(html_studio, "generate", gen)
    monkeypatch.setattr(html_studio.STATE, "seconds", 180.0)
    monkeypatch.setattr(page_shots, "shoot", shoot)
    monkeypatch.setattr(page_shots, "available", lambda: True)
    monkeypatch.setattr(weblibs, "ensure", no_download)
    monkeypatch.setattr(weblibs, "missing", lambda folder=None: [])
    return calls


async def test_make_polishes_by_screenshots(settings, polish_env):
    res = await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode())
    assert res.html == "<h1>v2</h1>" and res.cost_usd == 3.0
    assert polish_env["gen"] == [None, "<h1>v1</h1>"] and polish_env["shots"] == [b"<h1>v1</h1>", b"<h1>v2</h1>"]


async def test_make_keeps_draft_when_polish_breaks_or_fails(settings, polish_env):
    polish_env["after_errors"] = 2  # доводка добавила ошибок JS
    assert (await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode())).html == "<h1>v1</h1>"
    polish_env.update(after_errors=0, polish_fail=True, shots=[])
    res = await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode())
    assert res.html == "<h1>v1</h1>" and res.cost_usd == pytest.approx(2.7)  # и потраченное на доводку учтено
    polish_env.update(polish_fail=False, shots_fail=True, gen=[])
    assert (await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode())).html == "<h1>v1</h1>"
    assert polish_env["gen"] == [None]  # без скриншотов доводку не зовём
    settings.html_polish = False
    polish_env.update(shots_fail=False, shots=[], gen=[])
    await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode())
    assert polish_env["gen"] == [None] and polish_env["shots"] == []


def _browser_ready():
    from bot import page_shots

    if not page_shots.available():
        return False
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as p:
            import os

            return os.path.exists(p.chromium.executable_path)
    except Exception:
        return False


@pytest.mark.skipif(not _browser_ready(), reason="нет Playwright с Chromium")
async def test_real_screenshots_errors_and_no_network():
    from bot import page_shots

    page = media.html_page("t", '<div style="width:900px;height:50px;background:red"></div><button '
                                'onclick="document.body.style.background=\'blue\'">Играть</button>'
                                '<script>fetch("https://example.com").catch(()=>{}); null.x</script>').data
    shots = await page_shots.shoot(page)
    labels = [label for label, _ in shots.images]
    assert all(img[:2] == b"\xff\xd8" for _, img in shots.images)
    assert shots.overflow and shots.page_errors() == 1
    assert "через 2.5 с после нажатия на главную кнопку" in labels and labels[-1].startswith("маленький телефон")


@pytest.mark.skipif(not _browser_ready(), reason="нет Playwright с Chromium")
async def test_single_screen_game_shows_pause_and_final_screen():
    from bot import page_shots

    game = media.html_page("g", """<style>body{margin:0;height:100vh;overflow:hidden;background:#111;color:#fff}</style>
<button data-primary onclick="st.textContent='игра'">Играть</button>
<button data-pause onclick="st.textContent = st.textContent === 'пауза' ? 'игра' : 'пауза'">II</button>
<h1 id=st>старт</h1><script>window.__studioEnd = () => { st.textContent = 'финал'; };</script>""").data
    labels = [label for label, _ in (await page_shots.shoot(game)).images]
    assert "пауза (нажата кнопка data-pause)" in labels and "экран финала (window.__studioEnd)" in labels
    assert "после прокрутки вниз на экран" not in labels and len(labels) == 6


# ---------- в чате: задание → «делаю, минуты через 3» → файл ----------
@pytest.fixture
def studio(monkeypatch):
    calls = {"n": 0, "delay": 0.0, "fail": None, "args": []}

    async def run(settings, brief, title, image=None, build=None, kind="page"):
        calls["n"] += 1
        calls["args"].append((brief, title, image))
        await asyncio.sleep(calls["delay"])
        if calls["fail"]:
            raise html_studio.StudioError(calls["fail"], 0.25)
        html = PAGE.replace("Привет", "<img src='__PHOTO__'>Привет")
        calls["built"] = build(html)  # то же, что увидит браузер для скриншотов
        return html_studio.StudioResult(html, 0.5, 1.0)

    async def no_download(*a, **kw):
        return []

    monkeypatch.setattr(html_studio, "run", run)
    monkeypatch.setattr(weblibs, "ensure", no_download)
    monkeypatch.setattr(html_studio.STATE, "seconds", 180.0)
    html_studio.STATE.busy_chats.clear()
    yield calls
    html_studio.STATE.busy_chats.clear()


def _html_tool(reply, **extra):
    async def talk(tools, prompt):
        out = await tools["make_html"].handler({"title": "Звёзды", "brief": BRIEF, "caption": "играй!", **extra})
        assert "уже делается" in out and "Код писать не нужно" in out
        return reply
    return talk


async def feed(brain, *payloads):
    for p in payloads:
        await brain.handle_payload(p)
    await settle(brain, timeout=20)


async def _chat_id(brain):
    return (await brain.db.get_chat_by_jid("stella", GROUP))["id"]


async def test_reply_first_then_html_file_as_answer(brain, studio):
    brain.llm.queue = [_html_tool("делаю игру, минуты через 3 ✨")]
    await feed(brain, upsert("стелла сделай игру html", msg_id="ASK1"))
    assert [m["type"] for m in brain.evo.sent] == ["text", "media"]
    doc = brain.evo.sent[1]
    assert doc["mimetype"] == "text/html" and doc["filename"] == "Звёзды.html" and doc["caption"] == "играй!"
    assert doc["quoted"]["key"]["id"] == "ASK1"  # файл — ответом на просьбу
    assert studio["args"][0][:2] == (BRIEF, "Звёзды") and studio["args"][0][2] is None
    assert await brain.refresh_spend() >= 0.5  # расходы студии учтены в дневном бюджете


async def test_silent_model_still_announces_and_failure_refunds(brain, studio):
    brain.llm.queue = [_html_tool("<молчу>")]
    await feed(brain, upsert("стелла сайт"))
    assert brain.evo.texts()[0].startswith("делаю страницу ✨ будет готова")
    studio["fail"] = "страница делалась слишком долго"
    brain.llm.queue = [_html_tool("делаю")]
    await feed(brain, upsert("стелла ещё сайт"))
    assert brain.evo.texts()[-1] == "страница не получилась 😕 страница делалась слишком долго"
    used = await brain.db.kv_get(f"html:{await _chat_id(brain)}:{date.today().isoformat()}")
    assert used[1] == 1  # неудачная не съела лимит
    assert await brain.refresh_spend() >= 0.75  # а потраченное студией учтено


async def test_limits_busy_chat_and_bad_brief(brain, studio):
    errors = []

    async def try_tool(args):
        try:
            await brain_tools["make_html"].handler(args)
        except ToolError as exc:
            errors.append(str(exc))

    brain_tools = {}

    async def twice(tools, prompt):
        brain_tools.update(tools)
        await try_tool({"title": "x", "brief": "коротко"})
        await tools["make_html"].handler({"title": "x", "brief": BRIEF})
        await try_tool({"title": "y", "brief": BRIEF})
        return "делаю"

    brain.s.html_per_day = 1
    studio["delay"] = 0.3
    brain.llm.queue = [twice]
    await feed(brain, upsert("стелла сайт"))
    assert "опиши подробно" in errors[0] and errors[1] == "за один ответ — одна страница"

    async def again(tools, prompt):
        brain_tools.update(tools)
        await try_tool({"title": "x", "brief": BRIEF})
        return "ок"

    brain.llm.queue = [again]
    await feed(brain, upsert("стелла ещё"))
    assert "хватит" in errors[-1]
    brain.llm.queue = [_html_tool("ок")]
    await feed(brain, upsert("стелла сайт", sender=f"{OWNER}@s.whatsapp.net", name="Хозяин"))  # владельцу можно
    assert len([m for m in brain.evo.sent if m["type"] == "media"]) == 2


async def test_second_request_while_busy_is_refused(brain, studio):
    errors = []
    studio["delay"] = 0.5

    async def second(tools, prompt):
        try:
            await tools["make_html"].handler({"title": "x", "brief": BRIEF})
        except ToolError as exc:
            errors.append(str(exc))
        return "ок"

    brain.llm.queue = [_html_tool("делаю"), second]
    await brain.handle_payload(upsert("стелла сайт"))
    for _ in range(100):
        await asyncio.sleep(0.02)
        if brain.evo.texts():
            break
    await feed(brain, upsert("стелла и ещё один"))
    assert errors == ["я уже делаю страницу для этого чата — подожди, пока пришлю"]


async def test_budget_and_switch_off(brain, studio):
    errors = []

    async def grab(tools, prompt):
        try:
            await tools["make_html"].handler({"title": "x", "brief": BRIEF})
        except ToolError as exc:
            errors.append(str(exc))
        return "ок"

    await brain.db.kv_add(f"spend:{date.today().isoformat()}", brain.s.daily_budget_usd)
    brain.llm.queue = [grab]
    await feed(brain, upsert("стелла сайт"))
    brain.s.html_enabled = False
    brain.llm.queue = [grab]
    await feed(brain, upsert("стелла сайт"))
    assert "бюджет" in errors[0] and "выключены" in errors[1] and studio["n"] == 0


async def test_page_with_photo_from_chat(brain, studio):
    buf = io.BytesIO()
    Image.new("RGB", (64, 64), (10, 200, 10)).save(buf, "JPEG")
    brain.evo.media = buf.getvalue()
    brain.llm.queue = [_html_tool("делаю открытку", use_photo=True)]
    photo = {"stanzaId": "PH1", "participant": "77011111111@s.whatsapp.net",
             "quotedMessage": {"imageMessage": {"url": "x", "mimetype": "image/jpeg"}}}
    await feed(brain, upsert("стелла сделай открытку с этой фоткой", quoted=photo))
    assert studio["args"][0][2][0] == "image/jpeg"  # студия видит фото
    assert b"data:image/jpeg;base64," in studio["built"]  # и в браузер для скриншотов идёт страница с фото
    assert brain.evo.sent[-1]["type"] == "media" and brain.evo.sent[-1]["size"] > 1000  # фото внутри файла


async def test_photo_required_when_asked(brain, studio):
    errors = []

    async def ask(tools, prompt):
        try:
            await tools["make_html"].handler({"title": "x", "brief": BRIEF, "use_photo": True})
        except ToolError as exc:
            errors.append(str(exc))
        return "ок"

    brain.llm.queue = [ask]
    await feed(brain, upsert("стелла открытку с фото"))
    assert "фото нет" in errors[0]


async def test_paused_chat_gets_no_page(brain, studio):
    from datetime import timedelta

    from bot.util import utcnow

    studio["delay"] = 0.3

    async def talk(tools, prompt):
        await tools["make_html"].handler({"title": "x", "brief": BRIEF})
        await brain.db.update_chat(await _chat_id(brain), paused_until=utcnow() + timedelta(minutes=5))
        return "делаю"

    brain.llm.queue = [talk]
    await feed(brain, upsert("стелла сайт"))
    assert not [m for m in brain.evo.sent if m["type"] == "media"]
    used = await brain.db.kv_get(f"html:{await _chat_id(brain)}:{date.today().isoformat()}")
    assert not used or used[1] == 0


# ---------- презентации 10/10: слайды из студии -> PowerPoint + PDF ----------
DECK = """<!doctype html><html><head><style>
body{margin:0;background:#111} .slide{width:1920px;height:1080px;position:relative;overflow:hidden;
background:linear-gradient(135deg,#1b1640,#4a1d6b);color:#fff;font-family:'Unbounded',sans-serif}
h1{font-size:96px;margin:120px} .long{position:absolute;left:1700px;top:500px;font-size:48px;white-space:nowrap}
.tiny{font-size:12px;position:absolute;bottom:40px;left:120px}
</style></head><body>
<section class="slide"><h1>Марс</h1><aside class="notes">Начните с вопроса залу.</aside></section>
<section class="slide"><h1>Вода</h1><div class="long">Этот текст уезжает за край слайда</div></section>
<section class="slide"><h1>Итог</h1><p class="tiny">мелкая сноска</p></section>
</body></html>"""


def _jpeg(color=(40, 20, 90), size=(1920, 1080)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "JPEG")
    return buf.getvalue()


@pytest.mark.skipif(not _browser_ready(), reason="нет Playwright с Chromium")
async def test_real_deck_slides_notes_and_problems():
    from bot import page_shots

    shots = await page_shots.shoot_deck(media.html_page("Марс", DECK).data)
    assert len(shots.slides) == 3 and shots.slides[0].notes == "Начните с вопроса залу."
    assert Image.open(io.BytesIO(shots.slides[1].jpeg)).size == (1920, 1080)
    assert Image.open(io.BytesIO(shots.images[0][1])).size == (1600, 900) and shots.images[2][0] == "слайд 3"
    assert any(p.startswith("слайд 2: текст вылезает") for p in shots.problems)
    assert any(p.startswith("слайд 3: мелкий текст") for p in shots.problems)
    assert not any(p.startswith("слайд 1") for p in shots.problems) and shots.page_errors() == 0


def test_deck_files_pptx_with_notes_and_pdf():
    from pptx import Presentation

    from bot.page_shots import Slide

    pptx, pdf = media.deck_files("Марс: 10/10", [Slide(_jpeg(), "что сказать"), Slide(_jpeg((200, 50, 50)))])
    assert pptx.filename == "Марс 1010.pptx" and pptx.mimetype == media.MIME["pptx"] and pdf.filename.endswith(".pdf")
    prs = Presentation(io.BytesIO(pptx.data))
    assert len(prs.slides) == 2 and prs.slide_width == 12192000
    assert prs.slides[0].notes_slide.notes_text_frame.text == "что сказать"
    pic = prs.slides[0].shapes[0]
    assert (pic.left, pic.top, pic.width, pic.height) == (0, 0, prs.slide_width, prs.slide_height)
    assert pdf.data.startswith(b"%PDF") and pdf.data.count(b"/Type /Page\n") + pdf.data.count(b"/Type /Page ") >= 2
    with pytest.raises(media.MediaError):
        media.deck_files("x", [])


async def test_deck_generate_uses_deck_prompt_and_problems(settings, monkeypatch):
    from bot.page_shots import DeckShots

    fake = FakeClient(PAGE)
    monkeypatch.setattr(html_studio, "_client", fake)
    await html_studio.generate(settings, BRIEF, "Марс", kind="deck")
    assert fake.calls[0]["system"] == html_studio.DECK_SYSTEM and 'class="slide"' in html_studio.DECK_SYSTEM
    shots = DeckShots(images=[("слайд 1", b"j")], problems=["слайд 1: мелкий текст (меньше 20 px) — 2 шт."])
    await html_studio.generate(settings, BRIEF, "Марс", draft="<section class=slide></section>", shots=shots,
                               kind="deck")
    texts = "\n".join(b["text"] for b in fake.calls[1]["messages"][0]["content"] if b["type"] == "text")
    assert "Проверка нашла" in texts and "мелкий текст" in texts and "10 из 10" in texts


@pytest.fixture
def deck_env(monkeypatch, polish_env):
    from bot import page_shots
    from bot.page_shots import DeckShots, Slide

    state = {"slides_before": 3, "slides_after": 3}

    async def shoot_deck(page, timeout=0):
        polish_env["shots"].append(page)
        n = state["slides_before"] if len(polish_env["shots"]) == 1 else state["slides_after"]
        return DeckShots([Slide(b"j", f"n{i}") for i in range(n)], [("слайд", b"j")] * n)

    monkeypatch.setattr(page_shots, "shoot_deck", shoot_deck)
    return state


async def test_deck_make_renders_and_keeps_more_complete_version(settings, polish_env, deck_env):
    res = await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode(), kind="deck")
    assert res.html == "<h1>v2</h1>" and len(res.shots.slides) == 3
    polish_env["shots"].clear()
    deck_env["slides_after"] = 1  # доводка потеряла слайды — берём черновик
    res = await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode(), kind="deck")
    assert res.html == "<h1>v1</h1>" and len(res.shots.slides) == 3
    settings.html_polish = False  # без доводки слайды всё равно снимаются
    polish_env["shots"].clear()
    polish_env["gen"].clear()
    res = await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode(), kind="deck")
    assert polish_env["gen"] == [None] and len(res.shots.slides) == 3
    polish_env["shots"].clear()
    deck_env["slides_before"] = 0
    with pytest.raises(html_studio.StudioError, match="слайды"):
        await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode(), kind="deck")


async def test_deck_without_browser_fails_cleanly(settings, polish_env, monkeypatch):
    from bot import page_shots

    monkeypatch.setattr(page_shots, "available", lambda: False)
    with pytest.raises(html_studio.StudioError, match="нет браузера") as exc:
        await html_studio.make(settings, BRIEF, "x", None, build=lambda h: h.encode(), kind="deck")
    assert exc.value.cost_usd == 2.0


async def test_presentation_in_chat_arrives_as_pptx_and_pdf(brain, studio, monkeypatch):
    from bot import page_shots
    from bot.page_shots import DeckShots, Slide

    monkeypatch.setattr(page_shots, "available", lambda: True)

    async def run(settings, brief, title, image=None, build=None, kind="page"):
        studio["args"].append((brief, title, kind))
        return html_studio.StudioResult("<section class=slide></section>", 0.9, 1.0,
                                        shots=DeckShots([Slide(_jpeg(), "привет залу"), Slide(_jpeg())]))

    monkeypatch.setattr(html_studio, "run", run)

    async def talk(tools, prompt):
        assert "brief" in tools["make_presentation"].input_schema["properties"]
        out = await tools["make_presentation"].handler({"title": "Марс", "brief": "презентация про Марс для 7 класса, "
                                                        "12 слайдов, вода, роверы, колонизация", "caption": "держи"})
        assert "PowerPoint и PDF придут" in out and "Слайды писать не нужно" in out
        return "делаю презентацию, минуты через 3 🎨"

    brain.llm.queue = [talk]
    await feed(brain, upsert("стелла сделай презентацию про марс", msg_id="ASK7"))
    assert [m["type"] for m in brain.evo.sent] == ["text", "media", "media"]
    pptx, pdf = brain.evo.sent[1], brain.evo.sent[2]
    assert pptx["filename"] == "Марс.pptx" and pptx["caption"] == "держи" and pptx["quoted"]["key"]["id"] == "ASK7"
    assert pdf["filename"] == "Марс.pdf" and pdf["mimetype"] == media.MIME["pdf"]
    assert studio["args"][0][2] == "deck"


async def test_simple_presentation_without_browser(brain, monkeypatch):
    from bot import page_shots

    monkeypatch.setattr(page_shots, "available", lambda: False)
    seen = {}

    async def grab(tools, prompt):
        seen.update(tools)
        return "ок"

    brain.llm.queue = [grab]
    await feed(brain, upsert("стелла привет"))
    assert "slides" in seen["make_presentation"].input_schema["properties"]  # старая простая — из списка слайдов


# ---------- связь со студией: обрывы, перегрузка, понятные ошибки ----------
def _sdk_client(monkeypatch, replies):
    """Настоящий клиент Anthropic SDK поверх поддельного сервера: replies — по ответу на каждый запрос."""
    import json as _json

    import anthropic
    import httpx2

    monkeypatch.setattr(html_studio, "RETRY_PAUSE", 0)
    calls = []

    class Broken(httpx2.AsyncByteStream):
        async def __aiter__(self):
            yield (b'event: message_start\ndata: {"type":"message_start","message":{"id":"m","type":"message",'
                   b'"role":"assistant","model":"claude-opus-5-5","content":[],"stop_reason":null,'
                   b'"stop_sequence":null,"usage":{"input_tokens":5,"output_tokens":1}}}\n\n')
            raise httpx2.RemoteProtocolError("peer closed connection without sending complete message body")

    def ok():
        events = [
            {"type": "message_start", "message": {"id": "m", "type": "message", "role": "assistant",
                                                  "model": "claude-opus-5-5", "content": [], "stop_reason": None,
                                                  "stop_sequence": None,
                                                  "usage": {"input_tokens": 3000, "output_tokens": 1}}},
            {"type": "content_block_start", "index": 0, "content_block": {"type": "text", "text": ""}},
            {"type": "content_block_delta", "index": 0, "delta": {"type": "text_delta", "text": PAGE}},
            {"type": "content_block_stop", "index": 0},
            {"type": "message_delta", "delta": {"stop_reason": "end_turn", "stop_sequence": None},
             "usage": {"output_tokens": 20000}},
            {"type": "message_stop"},
        ]
        body = "".join(f"event: {e['type']}\ndata: {_json.dumps(e)}\n\n" for e in events).encode()
        return httpx2.Response(200, content=body, headers={"content-type": "text/event-stream"})

    async def handler(request):
        calls.append(request)
        kind = replies[min(len(calls), len(replies)) - 1]
        if kind == "ok":
            return ok()
        if kind == "drop":
            return httpx2.Response(200, stream=Broken(), headers={"content-type": "text/event-stream"})
        code, message = kind
        return httpx2.Response(code, json={"type": "error", "error": {"type": "x", "message": message}})

    client = anthropic.AsyncAnthropic(api_key="test", max_retries=0,
                                      http_client=httpx2.AsyncClient(transport=httpx2.MockTransport(handler)))
    monkeypatch.setattr(html_studio, "_client", client)
    return calls


async def test_dropped_stream_is_retried(settings, monkeypatch):
    calls = _sdk_client(monkeypatch, ["drop", "ok"])
    res = await html_studio.generate(settings, BRIEF, "x")
    assert res.html == PAGE and len(calls) == 2


async def test_dropped_twice_gives_clear_error(settings, monkeypatch):
    _sdk_client(monkeypatch, ["drop", "drop"])
    with pytest.raises(html_studio.StudioError, match="связь со студией оборвалась \\(RemoteProtocolError\\)"):
        await html_studio.generate(settings, BRIEF, "x")


async def test_overload_is_retried_and_api_errors_are_readable(settings, monkeypatch):
    calls = _sdk_client(monkeypatch, [(529, "Overloaded"), "ok"])
    assert (await html_studio.generate(settings, BRIEF, "x")).html == PAGE and len(calls) == 2
    _sdk_client(monkeypatch, [(400, "Your credit balance is too low to access the Anthropic API.")])
    with pytest.raises(html_studio.StudioError, match="кончились деньги"):
        await html_studio.generate(settings, BRIEF, "x")
    _sdk_client(monkeypatch, [(404, "model: claude-opus-5-5")])
    with pytest.raises(html_studio.StudioError, match="HTML_MODEL=claude-sonnet-5-5"):
        await html_studio.generate(settings, BRIEF, "x")
    calls = _sdk_client(monkeypatch, [(401, "invalid x-api-key")])
    with pytest.raises(html_studio.StudioError, match="ANTHROPIC_API_KEY"):
        await html_studio.generate(settings, BRIEF, "x")
    assert len(calls) == 1  # неверный ключ не повторяем


async def test_max_tokens_retries_one_effort_lower(settings, monkeypatch):
    fake = FakeClient(replies=[("<!doctype html><html><body>обрыв", "max_tokens"), (PAGE, "end_turn")])
    monkeypatch.setattr(html_studio, "_client", fake)
    res = await html_studio.generate(settings, BRIEF, "x")
    assert res.html == PAGE and res.cost_usd == pytest.approx(2 * 0.412)
    assert [c["output_config"]["effort"] for c in fake.calls] == ["max", "high"]
    # документ целиком уже написан, а дальше модель продолжала болтать — берём его
    fake = FakeClient(replies=[(PAGE + "\nи ещё пояснения…", "max_tokens")])
    monkeypatch.setattr(html_studio, "_client", fake)
    assert (await html_studio.generate(settings, BRIEF, "x")).html == PAGE and len(fake.calls) == 1
    settings.html_effort = "high"  # ниже high не опускаемся
    fake = FakeClient(replies=[("<html>обрыв", "max_tokens")])
    monkeypatch.setattr(html_studio, "_client", fake)
    with pytest.raises(html_studio.StudioError, match="слишком большой"):
        await html_studio.generate(settings, BRIEF, "x")
    assert len(fake.calls) == 1


async def test_budget_rejected_by_api_falls_back(settings, monkeypatch):
    calls = _sdk_client(monkeypatch, [(400, "output_config.task_budget: unknown field"), "ok"])
    res = await html_studio.generate(settings, BRIEF, "x")
    assert res.html == PAGE and len(calls) == 2
    import json as _json

    first, second = (_json.loads(c.content) for c in calls)
    assert "task_budget" in first["output_config"] and "task_budget" not in second["output_config"]
    assert calls[0].headers.get("anthropic-beta") == "task-budgets-2026-03-13"
    assert "task-budgets" not in (calls[1].headers.get("anthropic-beta") or "")


# ---------- после ревью кода ----------
def test_google_fonts_import_with_semicolons_is_removed_cleanly(tmp_path):
    css = ("<style>@import url('https://fonts.googleapis.com/css2?family=Unbounded:wght@400;700;900&family=Inter:"
           "wght@400;600&display=swap');\n:root{--bg:#0b0b1a;--accent:#ff3d7f}\nh1{font-family:'Unbounded'}</style>")
    text = media.html_page("t", css, libs_dir=str(tmp_path)).data.decode()
    assert "googleapis" not in text and "700;900" not in text and ":root{--bg:#0b0b1a;--accent:#ff3d7f}" in text
    plain = media.html_page("t", '<style>@import "https://fonts.googleapis.com/css?family=Lobster";\nb{x:1}</style>',
                            libs_dir=str(tmp_path)).data.decode()
    assert "googleapis" not in plain and "b{x:1}" in plain


def test_extract_html_keeps_pages_with_their_own_code_fences():
    cheat = ('<!doctype html><html><body><h1>Markdown</h1><pre>```\n# Заголовок\n```</pre>'
             '<pre>```js\n<div>const</div> x = 1;\n```</pre></body></html>')
    assert html_studio.extract_html(cheat) == cheat
    assert html_studio.extract_html(f"Вот:\n```html\n{cheat}\n```\nготово") == cheat
    fragment = "<style>b{}</style><div>```\nкод\n```</div>"
    assert html_studio.extract_html(f"```html\n{fragment}\n```") == fragment


def test_regexes_stay_fast_on_hostile_html(tmp_path):
    import time

    for junk in ("<link " * 60000, "<meta " * 60000, "@import url(" * 30000, "<script " * 60000, "<base " * 60000):
        started = time.monotonic()
        media.html_page("t", junk[: media.MAX_HTML_CHARS], libs_dir=str(tmp_path))
        assert time.monotonic() - started < 2, junk[:10]


def test_resource_hints_are_stripped_and_dns_prefetch_off(tmp_path):
    text = media.html_page("t", '<link rel="dns-prefetch" href="//evil.example"><link rel=preconnect href="https://x">'
                                '<link rel="icon" href="data:,"><h1>x</h1>', libs_dir=str(tmp_path)).data.decode()
    assert "evil.example" not in text and "preconnect" not in text and 'rel="icon"' in text
    assert 'http-equiv="x-dns-prefetch-control" content="off"' in text and "webrtc" not in text


def test_adaptive_thinking_only_for_models_that_take_it():
    yes = ("claude-opus-5-5", "claude-opus-5", "claude-sonnet-5-5", "claude-haiku-5-5", "claude-opus-4-8",
           "claude-opus-4-6", "claude-sonnet-4-6", "claude-fable-5-1", "claude-mythos-5")
    no = ("claude-opus-4-5", "claude-sonnet-4-5", "claude-haiku-4-5", "claude-sonnet-4-20250514",
          "claude-opus-4-1", "some-other-model")
    assert all(html_studio._thinking(m) for m in yes), [m for m in yes if not html_studio._thinking(m)]
    assert not any(html_studio._thinking(m) for m in no), [m for m in no if html_studio._thinking(m)]


def test_browser_folder_is_passed_to_the_worker(monkeypatch, tmp_path):
    from bot import page_shots

    monkeypatch.delenv("PLAYWRIGHT_BROWSERS_PATH", raising=False)
    monkeypatch.setenv("HOME", str(tmp_path))
    (tmp_path / ".cache" / "ms-playwright" / "chromium_headless_shell-1").mkdir(parents=True)
    env = page_shots._env("/tmp/worker-home")
    assert env["PLAYWRIGHT_BROWSERS_PATH"] == str(tmp_path / ".cache" / "ms-playwright") and env["HOME"] != str(tmp_path)
    assert "ANTHROPIC_API_KEY" not in env


def test_partial_cost_of_a_cut_stream_is_estimated():
    import time

    stream = SimpleNamespace(current_message_snapshot=SimpleNamespace(usage=SimpleNamespace(input_tokens=4000)))
    base = {"model": "claude-opus-5-5", "max_tokens": 128000}
    cost = html_studio._partial_cost(base, {"stream": stream, "started": time.monotonic() - 100})
    assert 0.15 < cost < 0.2  # ~8000 токенов за 100 с + вход
    capped = html_studio._partial_cost(base, {"stream": stream, "started": time.monotonic() - 1e5})
    assert capped == pytest.approx((4000 * 4 + 128000 * 20) / 1e6)
    assert html_studio._partial_cost(base, {}) == 0.0


def test_polish_that_adds_sideways_scroll_is_rejected():
    from bot.page_shots import Shots

    assert html_studio._worse(Shots(overflow=True), Shots(overflow=False), deck=False)
    assert not html_studio._worse(Shots(overflow=True), Shots(overflow=True), deck=False)


async def test_budget_counts_pages_still_in_progress(brain, studio, monkeypatch):
    errors = []

    async def grab(tools, prompt):
        try:
            await tools["make_html"].handler({"title": "x", "brief": BRIEF})
        except ToolError as exc:
            errors.append(str(exc))
        return "ок"

    await brain.db.kv_add(f"spend:{date.today().isoformat()}", brain.s.daily_budget_usd - 1.0)
    monkeypatch.setattr(html_studio.STATE, "reserved_usd", 1.5)  # ещё делается страница на ~1.5 $
    brain.llm.queue = [grab]
    await feed(brain, upsert("стелла сайт"))
    assert errors and "бюджет" in errors[0] and studio["n"] == 0


@pytest.mark.skipif(not _browser_ready(), reason="нет Playwright с Chromium")
async def test_deck_checks_find_what_the_eye_sees():
    from bot import page_shots

    deck = """<!doctype html><html><head><style>
body{margin:0} .slide{width:1920px;height:1080px;position:relative;overflow:hidden;background:#141428;color:#fff}
h1{font-size:80px;margin:0;position:absolute;left:120px;top:150px;width:1500px}
.card{position:absolute;left:120px;top:420px;width:300px;height:80px;overflow:hidden;background:#223;font-size:40px}
.a{position:absolute;left:120px;top:700px;font-size:48px}.b{position:absolute;left:180px;top:705px;font-size:48px}
.blob{position:absolute;left:900px;top:380px;width:400px;height:200px;background:#e55}
.under{position:absolute;left:900px;top:430px;font-size:40px}
.foot{position:absolute;bottom:40px;left:120px;font-size:18px}.empty{position:absolute;inset:300px 200px}
</style></head><body>
<svg width="0" height="0" style="position:absolute"><symbol id="logo" viewBox="-50 -50 100 100">
  <text x="-49" y="10">Go</text></symbol></svg>
<section class="slide"><h1>Чистый слайд — в 2021 году было 6 779 км</h1>
  <div class="foot" data-chrome>Марс · 01 / 03</div><div class="empty"></div>
  <svg style="position:absolute;left:0;top:0" width="1920" height="1080"><circle cx="1500" cy="700" r="200"
    fill="none" stroke="#fff" stroke-width="4"/></svg>
  <p style="position:absolute;left:1420px;top:680px;font-size:40px;margin:0">в круге</p>
  <svg style="position:absolute;left:300px;top:600px" width="100" height="100"><use href="#logo"/></svg></section>
<section class="slide"><h1>Дефекты</h1>
  <div class="card">Этот длинный текст в карточке обрезается снизу и справа точно</div>
  <div class="a">Наложение</div><div class="b">другой текст</div>
  <div class="under">текст под блобом</div><div class="blob"></div>
  <svg style="position:absolute;left:0;top:0" width="1920" height="1080"><line x1="100" y1="884" x2="1800" y2="884"
    stroke="#fff" stroke-width="3"/></svg>
  <p style="position:absolute;left:300px;top:860px;font-size:40px;z-index:2;margin:0">линия под текстом</p>
  <p style="position:absolute;left:1700px;top:250px;font-size:40px;white-space:nowrap">уезжает за край</p>
  <p style="position:absolute;left:120px;top:1000px;font-size:14px">крошечная сноска</p></section>
<section class="slide"><p class="logo" style="position:absolute;left:600px;top:400px;font-size:120px;margin:0">ШаурмаGo</p>
  <svg style="position:absolute;left:0;top:0;z-index:3" width="1920" height="1080"><circle cx="1180" cy="470" r="120"
    fill="none" stroke="#f80" stroke-width="34"/></svg>
  <svg style="position:absolute;left:1500px;top:300px" width="380" height="380"><rect width="380" height="380"
    fill="#0af"/></svg></section>
</body></html>"""
    shots = await page_shots.shoot_deck(media.html_page("t", deck).data)
    text = "\n".join(shots.problems)
    assert not any(p.startswith("слайд 1") for p in shots.problems), text  # без ложных тревог (и текст в <symbol>)
    for expected in ("слайд 2: текст вылезает за край слайда («уезжает за край» в <p>)", "обрезан рамкой карточки",
                     "«Наложение» в <div.a> и «другой текст» в <div.b>", "перекрыт рисунком или декором («текст под "
                     "блобом» в <div.under>)", "линия или кольцо проходит под текстом («линия под текстом» в <p>)",
                     "мелкий текст", "слайд 3: текст перекрыт рисунком или декором («ШаурмаGo» в <p.logo>)",
                     "слайд 3: рисунок почти у края слайда (<svg> 40 px)"):
        assert expected in text, (expected, text)
    assert shots.geometry == 5


def test_ruble_sign_brings_the_latin_ext_font_subset(tmp_path, monkeypatch):
    cyr = _asset(tmp_path, "inter-cyrillic.woff2", "x@1/c.woff2", b"cyr")
    lat = _asset(tmp_path, "inter-latin.woff2", "x@1/l.woff2", b"lat")
    ext = _asset(tmp_path, "inter-latin-ext.woff2", "x@1/e.woff2", b"ext-with-ruble")
    monkeypatch.setitem(weblibs.FONTS, "Inter", ("100 900", (cyr, lat, ext)))
    import base64

    marker = base64.b64encode(b"ext-with-ruble").decode()
    with_rub = media.html_page("t", "<p style='font-family:Inter'>Средний чек 450 ₽</p>", libs_dir=str(tmp_path))
    plain = media.html_page("t", "<p style='font-family:Inter'>Средний чек 450 рублей</p>", libs_dir=str(tmp_path))
    assert marker in with_rub.data.decode() and "U+20AD-20C0" in with_rub.data.decode()
    assert marker not in plain.data.decode()
    assert all(len(assets) == 3 for _, assets in weblibs.FONTS.values())


@pytest.mark.skipif(not _browser_ready(), reason="нет Playwright с Chromium")
async def test_typography_keeps_code_and_does_not_glue_numbers():
    from playwright.async_api import async_playwright

    from bot import page_shots

    html = ('<section class="slide"><pre id="code">x = 5\nfor i in range(3):\n    print(i)</pre>'
            '<p id="t">iPhone 6 128 ГБ и Model 3 500 км, 12 742 км — это много</p>'
            '<p id="poem" style="white-space:pre-line">ул. Тверская, 13\nТелефон</p></section>')
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page()
        await page.set_content(html)
        await page.evaluate(page_shots.DECK_TYPO_JS)
        code = await page.evaluate("document.getElementById('code').textContent")
        text = await page.evaluate("document.getElementById('t').textContent")
        poem = await page.evaluate("document.getElementById('poem').textContent")
        await browser.close()
    assert code == "x = 5\nfor i in range(3):\n    print(i)" and poem == "ул. Тверская, 13\nТелефон"
    assert "6 128" in text and "3 500" in text and "6128" not in text and "12 742" in text
    assert " — это" in text


@pytest.mark.skipif(not _browser_ready(), reason="нет Playwright с Chromium")
async def test_dense_headings_and_big_numbers_are_not_overlaps():
    from bot import page_shots

    deck = """<style>body{margin:0}.slide{width:1920px;height:1080px;position:relative;overflow:hidden;background:#111;
color:#fff;font-family:'Unbounded'}</style><section class="slide">
<h1 style="position:absolute;left:120px;top:150px;font-size:96px;line-height:1.05;margin:0">
<span style="display:block">Марс потерял</span><span style="display:block;color:#f84">свой океан</span></h1>
<div style="position:absolute;left:120px;top:520px"><div style="font-size:180px;line-height:1">24 ч</div>
<div style="font-size:32px;font-family:'Inter'">сутки на Марсе длиннее земных</div></div></section>"""
    shots = await page_shots.shoot_deck(media.html_page("t", deck).data)
    assert shots.geometry == 0 and not shots.problems, shots.problems


@pytest.mark.skipif(not _browser_ready(), reason="нет Playwright с Chromium")
async def test_hanging_final_hook_and_inner_scroll_do_not_break_shots():
    import time

    from bot import page_shots

    game = media.html_page("g", """<style>body{margin:0;height:100vh;overflow:hidden}</style>
<button data-primary>Играть</button><button data-pause>II</button>
<script>window.__studioEnd = () => new Promise(() => {});</script>""").data
    started = time.monotonic()
    labels = [label for label, _ in (await page_shots.shoot(game)).images]
    assert time.monotonic() - started < 60 and "пауза (нажата кнопка data-pause)" in labels
    assert labels[-1].startswith("маленький телефон")
    site = media.html_page("s", """<style>html,body{margin:0;height:100%;overflow:hidden}
main{height:100dvh;overflow-y:auto}section{height:900px}</style><main><section>один</section>
<section>два</section></main>""").data
    labels = [label for label, _ in (await page_shots.shoot(site)).images]
    assert "после прокрутки вниз на экран" in labels
